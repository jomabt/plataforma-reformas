"""PRO-REFORMAS: cálculo técnico orientativo, materiales, presupuestos y documentos
(versión en un solo archivo: Excel, Word y PDF con encabezado, y biblioteca de normas).

AVISO: resultados ORIENTATIVOS. No sustituyen la memoria técnica ni el proyecto de un
instalador autorizado o técnico competente.
"""
from __future__ import annotations

import base64
import hashlib
import math
import re
from collections import Counter
from datetime import date
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape as xml_escape

import pandas as pd
import streamlit as st

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.properties import PageSetupProperties
    HAY_OPENPYXL = True
except ImportError:
    HAY_OPENPYXL = False

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    HAY_DOCX = True
except ImportError:
    HAY_DOCX = False

try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.pdfgen import canvas as rl_canvas
    from reportlab.platypus import (CondPageBreak, PageBreak, Paragraph, SimpleDocTemplate,
                                    Spacer, Table, TableStyle)
    HAY_REPORTLAB = True
except ImportError:
    HAY_REPORTLAB = False

# Estado de la normativa que usa la app. Actualízalo cada vez que revises el BOE.
NORMATIVA_FECHA = "octubre de 2026"
NORMATIVA_APLICADA = "REBT (RD 842/2002) · CTE DB-HS4 · RITE (RD 1027/2007)"


# ----------------------------------------------------------------------------
# ELECTRICIDAD (REBT)
# ----------------------------------------------------------------------------
TENSION_V = 230.0
# Conductividad del cobre a temperatura de servicio (no a 20 °C, que da 56).
CONDUCTIVIDAD_CU = {"PVC (70 °C)": 48.0, "XLPE/EPR (90 °C)": 44.0}
SECCIONES_MM2 = [1.5, 2.5, 4, 6, 10, 16, 25, 35, 50]
# Intensidad admisible ORIENTATIVA (Cu, empotrado en tubo). Valores conservadores (PVC):
# con XLPE la admisible real es mayor. Verificar con ITC-BT-19, tabla 1.
I_ADMISIBLE_A = {1.5: 15, 2.5: 21, 4: 27, 6: 36, 10: 50, 16: 66, 25: 84, 35: 104, 50: 125}
SECCION_MIN_DI_MM2 = 6  # ITC-BT-15

# Caída de tensión máxima en la derivación individual según el tipo de suministro (ITC-BT-15).
CAIDA_DI_PCT = {
    "Contadores totalmente concentrados (1 %)": 1.0,
    "Contadores en más de un lugar (0,5 %)": 0.5,
    "Usuario único, sin línea general (1,5 %)": 1.5,
}
SUMINISTRO_DI_DEFECTO = "Contadores totalmente concentrados (1 %)"

POTENCIAS_PREVISTAS_W = {"Básica": [5750], "Elevada": [9200, 11500, 14490]}
IGA_ESTANDAR_A = [10, 16, 20, 25, 32, 40, 50, 63]

# Circuitos (ITC-BT-25): descripción, PIA (A), sección (mm²), tubo (mm)
CIRCUITOS = {
    "C1": ("Iluminación", 10, 1.5, 16),
    "C2": ("Tomas de uso general y frigorífico", 16, 2.5, 20),
    "C3": ("Cocina y horno", 25, 6, 25),
    "C4": ("Lavadora, lavavajillas y termo", 20, 4, 20),
    "C5": ("Baños y tomas auxiliares de cocina", 16, 2.5, 20),
    "C8": ("Calefacción eléctrica", 25, 6, 25),
    "C9": ("Aire acondicionado", 25, 6, 25),
    "C10": ("Secadora", 16, 2.5, 20),
}
MAX_PUNTOS = {"C1": 30, "C2": 20, "C5": 6}

# Mínimos SIMPLIFICADOS por estancia: (puntos de luz, tomas C2, tomas C5).
# En salones (si se indica su superficie) y pasillos se ajustan por superficie/longitud
# (ITC-BT-25, tabla 2).
MINIMOS_ESTANCIA = {
    "vestibulos": (1, 1, 0),
    "dormitorios": (1, 3, 0),
    "salones": (1, 3, 0),
    "banos": (1, 0, 1),
    "aseos": (1, 0, 1),
    "cocinas": (1, 2, 3),
    "galerias": (1, 1, 0),
    "terrazas": (1, 0, 0),
    "trasteros": (1, 1, 0),
}


def grado_recomendado(superficie_m2: float, calefaccion: bool, aire: bool) -> str:
    """Elevada si superficie > 160 m², calefacción eléctrica o aire acondicionado."""
    return "Elevada" if (superficie_m2 > 160 or calefaccion or aire) else "Básica"


def derivacion_individual(potencia_w: float, longitud_m: float,
                          aislamiento: str = "PVC (70 °C)", cos_phi: float = 1.0,
                          caida_max_pct: float = 1.0) -> dict:
    """Sección de la derivación individual (monofásica): máximo entre caída de tensión,
    intensidad admisible y el mínimo reglamentario de 6 mm²."""
    if potencia_w <= 0 or longitud_m <= 0:
        raise ValueError("Potencia y longitud deben ser mayores que 0.")
    gamma = CONDUCTIVIDAD_CU[aislamiento]
    intensidad = potencia_w / (TENSION_V * cos_phi)
    e_max = TENSION_V * caida_max_pct / 100
    s_caida = 2 * longitud_m * potencia_w / (gamma * e_max * TENSION_V)
    s_int = next((s for s in SECCIONES_MM2 if I_ADMISIBLE_A[s] >= intensidad), None)
    if s_int is None:
        raise ValueError("Intensidad fuera de tabla: calcular manualmente.")
    minimo = max(SECCION_MIN_DI_MM2, s_caida, s_int)
    s_final = next((s for s in SECCIONES_MM2 if s >= minimo), None)
    if s_final is None:
        raise ValueError("Sección > 50 mm²: revisar longitud/potencia o usar otra solución.")
    caida_pct = (2 * longitud_m * potencia_w / (gamma * s_final * TENSION_V)) / TENSION_V * 100
    return {"intensidad_a": intensidad, "s_caida": s_caida, "s_int": s_int,
            "seccion_mm2": s_final, "caida_pct": caida_pct,
            "caida_max_pct": caida_max_pct, "cumple": caida_pct <= caida_max_pct + 1e-9}


def circuitos(estancias: dict, pasillo_m: float = 0.0, extras: dict | None = None,
              sup_salon_m2: float = 0.0, superficie_m2: float = 0.0, grado: str = "Básica"):
    """Circuitos y puntos mínimos. Devuelve (filas_tabla, resumen_puntos)."""
    extras = extras or {}
    luz = c2 = c5 = 0
    for tipo, n in estancias.items():
        if tipo == "salones" and sup_salon_m2 > 0:
            # Salón: 1 punto de luz (2 si > 10 m²); 1 toma por cada 6 m² (mínimo 3).
            luz += n * (1 if sup_salon_m2 <= 10 else 2)
            c2 += n * max(3, math.ceil(sup_salon_m2 / 6))
            continue
        l, t2, t5 = MINIMOS_ESTANCIA[tipo]
        luz, c2, c5 = luz + l * n, c2 + t2 * n, c5 + t5 * n
    if pasillo_m > 0:  # 1 punto de luz y 1 toma por cada 5 m de pasillo
        luz += math.ceil(pasillo_m / 5)
        c2 += math.ceil(pasillo_m / 5)

    n_circ = {
        "C1": math.ceil(luz / MAX_PUNTOS["C1"]) if luz else 0,
        "C2": math.ceil(c2 / MAX_PUNTOS["C2"]) if c2 else 0,
        "C3": 1 if estancias.get("cocinas", 0) > 0 else 0,
        "C4": 1 if (estancias.get("cocinas", 0) > 0 or estancias.get("galerias", 0) > 0) else 0,
        "C5": math.ceil(c5 / MAX_PUNTOS["C5"]) if c5 else 0,
        "C8": 1 if extras.get("C8") else 0,
        "C9": 1 if extras.get("C9") else 0,
        "C10": 1 if extras.get("C10") else 0,
    }
    # Electrificación elevada con más de 160 m²: circuito C2 adicional (C7).
    c7_por_superficie = False
    if grado == "Elevada" and superficie_m2 > 160 and 0 < n_circ["C2"] < 2:
        n_circ["C2"] = 2
        c7_por_superficie = True

    puntos = {"C1": luz, "C2": c2, "C5": c5}
    filas = []
    for cod, n in n_circ.items():
        if n:
            desc, pia, seccion, tubo = CIRCUITOS[cod]
            filas.append({"Circuito": cod, "Descripción": desc, "Nº circuitos": n,
                          "PIA (A)": pia, "Sección (mm²)": seccion, "Tubo (mm)": tubo,
                          "Puntos previstos": puntos.get(cod, "—")})
    return filas, {"puntos_luz": luz, "tomas_c2": c2, "tomas_c5": c5,
                   "c7_por_superficie": c7_por_superficie}


def protecciones(potencia_w: float, n_circuitos: int, con_spd: bool) -> list[dict]:
    """Protecciones orientativas del cuadro: IGA, diferenciales, PIA y sobretensiones."""
    intensidad = potencia_w / TENSION_V
    iga = next((a for a in IGA_ESTANDAR_A if a >= intensidad - 1e-6), None)
    n_dif = math.ceil(n_circuitos / 5) if n_circuitos else 0
    return [
        {"Elemento": "IGA (magnetotérmico general)",
         "Valor": f"{iga} A" if iga else "Revisar (fuera de tabla)",
         "Observaciones": f"Para {potencia_w:.0f} W a 230 V ({intensidad:.1f} A)"},
        {"Elemento": "Diferenciales 30 mA", "Valor": f"{n_dif} ud",
         "Observaciones": "Máximo 5 circuitos por diferencial; intensidad ≥ 25 A; tipo A "
                          "recomendado (el tipo AC dejaría de ser válido con el nuevo REBT, "
                          "aún sin publicar)"},
        {"Elemento": "PIA por circuito", "Valor": f"{n_circuitos} ud",
         "Observaciones": "Calibres en la tabla de circuitos"},
        {"Elemento": "Protector contra sobretensiones",
         "Valor": "Incluido" if con_spd else "No incluido",
         "Observaciones": "Hoy depende del tipo de instalación y del riesgo (ITC-BT-23); el "
                          "borrador del nuevo REBT lo haría obligatorio"},
    ]


# ----------------------------------------------------------------------------
# FONTANERÍA (CTE DB-HS4)
# ----------------------------------------------------------------------------
# Caudal instantáneo mínimo por aparato (dm³/s): (agua fría, ACS). CTE DB-HS4, tabla 2.1.
APARATOS = {
    "lavabo": (0.10, 0.065), "ducha": (0.20, 0.10), "inodoro": (0.10, 0.0),
    "fregadero": (0.20, 0.10), "lavavajillas": (0.15, 0.10), "lavadora": (0.20, 0.15),
}


def aparatos_contados(n_banos: int, n_aseos: int, n_cocinas: int, n_galerias: int) -> Counter:
    ap = Counter()
    ap["lavabo"] = n_banos + n_aseos
    ap["ducha"] = n_banos
    ap["inodoro"] = n_banos + n_aseos
    ap["fregadero"] = n_cocinas
    ap["lavavajillas"] = n_cocinas
    ap["lavadora"] = n_galerias if n_galerias else min(n_cocinas, 1)
    return ap


def hidraulica(n_banos: int, n_aseos: int, n_cocinas: int, n_galerias: int,
               velocidad_ms: float = 1.5) -> dict:
    """Caudal de diseño con coeficiente de simultaneidad 1/√(n−1) y diámetro interior mínimo."""
    ap = aparatos_contados(n_banos, n_aseos, n_cocinas, n_galerias)
    n = sum(ap.values())
    q_fria = sum(c * APARATOS[a][0] for a, c in ap.items())
    q_acs = sum(c * APARATOS[a][1] for a, c in ap.items())
    kn = 1.0 if n <= 2 else 1 / math.sqrt(n - 1)
    q_diseno = q_fria * kn
    diametro = math.sqrt(4 * (q_diseno / 1000) / (math.pi * velocidad_ms)) * 1000 if q_diseno else 0.0
    return {"n_aparatos": n, "q_fria": q_fria, "q_acs": q_acs, "kn": kn,
            "q_diseno": q_diseno, "diametro_mm": diametro}


# ----------------------------------------------------------------------------
# CLIMATIZACIÓN (estimación por ratio; NO es un cálculo de cargas RITE)
# ----------------------------------------------------------------------------
RATIOS_W_M2 = {"Buena (80 W/m²)": 80, "Media (100 W/m²)": 100, "Deficiente (120 W/m²)": 120}


def climatizacion(superficie_m2: float, w_m2: float) -> dict:
    kw = superficie_m2 * w_m2 / 1000
    return {"kw": kw, "kcal_h": kw * 860}


# ----------------------------------------------------------------------------
# CAPÍTULOS DEL PROYECTO
# ----------------------------------------------------------------------------
CAP_ELEC = "Electricidad"
CAP_FONT = "Fontanería"
CAP_TELE = "Telecomunicaciones"
CAP_CCTV = "CCTV y videovigilancia"
CAP_CLIMA = "Climatización"
# (capítulo, clave de inclusión en el estado, nombre en los documentos)
CAPITULOS = [
    (CAP_ELEC, "inc_elec", "Instalación eléctrica y CGMP (REBT)"),
    (CAP_FONT, "inc_font", "Fontanería y evacuación (CTE DB-HS)"),
    (CAP_TELE, "inc_teleco", "Telecomunicaciones y red de datos"),
    (CAP_CCTV, "inc_cctv", "CCTV y videovigilancia"),
    (CAP_CLIMA, "inc_clima", "Climatización"),
]
NOMBRE_CAP = {cap: nombre for cap, _, nombre in CAPITULOS}


# ----------------------------------------------------------------------------
# CCTV / VIDEOVIGILANCIA (dimensionado orientativo)
# ----------------------------------------------------------------------------
CCTV_TECNOLOGIAS = ["IP (PoE)", "Analógica HD (HD-TVI/CVI)"]
# Caudal de grabación aproximado por cámara (Mbps, códec H.265, grabación continua).
CCTV_RESOLUCIONES = {"2 MP (1080p) · ≈ 2 Mbps": 2.0, "4 MP · ≈ 4 Mbps": 4.0,
                     "8 MP (4K) · ≈ 8 Mbps": 8.0}
CANALES_GRABADOR = [4, 8, 16, 32]
DISCOS_TB = [2, 4, 6, 8, 12, 16]
PUERTOS_POE = [4, 8, 16, 24]
W_POE_FIJA, W_POE_PTZ = 7.0, 25.0   # consumo orientativo por cámara (W)
MARGEN_ALMACENAMIENTO = 1.10        # 10 % de margen sobre el cálculo


def cctv(n_domo: int, n_bullet: int, n_ptz: int, tecnologia: str, mbps: float,
         dias: int, dist_m: float, monitor: bool = True) -> dict:
    """Dimensionado orientativo: grabador, almacenamiento, PoE y cableado."""
    n_fijas = n_domo + n_bullet
    n_cam = n_fijas + n_ptz
    ip = tecnologia == CCTV_TECNOLOGIAS[0]
    base = {"n_domo": n_domo, "n_bullet": n_bullet, "n_ptz": n_ptz, "n_fijas": n_fijas,
            "n_cam": n_cam, "ip": ip, "tecnologia": tecnologia, "mbps": mbps, "dias": dias,
            "dist_m": dist_m, "monitor": monitor}
    if n_cam <= 0:
        return base
    canales = next((c for c in CANALES_GRABADOR if c >= n_cam), CANALES_GRABADOR[-1])
    n_grab = math.ceil(n_cam / canales)
    gb_dia = n_cam * mbps * 10.8                      # 1 Mbps continuo ≈ 10,8 GB/día
    tb = gb_dia * dias * MARGEN_ALMACENAMIENTO / 1000
    if tb <= DISCOS_TB[-1]:
        disco, n_discos = next(d for d in DISCOS_TB if d >= tb), 1
    else:
        disco, n_discos = DISCOS_TB[-1], math.ceil(tb / DISCOS_TB[-1])
    poe_w = n_fijas * W_POE_FIJA + n_ptz * W_POE_PTZ if ip else 0.0
    puertos = next((p for p in PUERTOS_POE if p >= n_cam), PUERTOS_POE[-1])
    base.update({
        "canales": canales, "n_grabadores": n_grab, "gb_dia": gb_dia, "tb": tb,
        "disco_tb": disco, "n_discos": n_discos,
        "poe_w": poe_w, "poe_total_w": poe_w * 1.25,
        "puertos_poe": puertos, "n_switch": math.ceil(n_cam / puertos) if ip else 0,
        "fuentes": 0 if ip else math.ceil(n_cam / 8),
        "cable_m": math.ceil(n_cam * dist_m)})
    return base


# ----------------------------------------------------------------------------
# CATÁLOGO DE MATERIALES (precios propios: todos parten de 0 €)
# ----------------------------------------------------------------------------
CATALOGO: list[tuple[str, str, str, str]] = []   # (clave, capítulo, material, unidad)


def _cat(cap, filas):
    for clave, nombre, ud in filas:
        CATALOGO.append((clave, cap, nombre, ud))


_cat(CAP_ELEC, [
    ("el_di_6", "Cable unipolar ES07Z1-K libre de halógenos 6 mm² (derivación individual)", "m"),
    ("el_di_10", "Cable unipolar ES07Z1-K libre de halógenos 10 mm² (derivación individual)", "m"),
    ("el_di_16", "Cable unipolar ES07Z1-K libre de halógenos 16 mm² (derivación individual)", "m"),
    ("el_di_25", "Cable unipolar ES07Z1-K libre de halógenos 25 mm² (derivación individual)", "m"),
    ("el_di_35", "Cable unipolar ES07Z1-K libre de halógenos 35 mm² (derivación individual)", "m"),
    ("el_di_50", "Cable unipolar ES07Z1-K libre de halógenos 50 mm² (derivación individual)", "m"),
    ("el_tubo_di", "Tubo para derivación individual (diámetro según ITC-BT-15)", "m"),
    ("el_cable_15", "Cable H07V-K / ES07Z1-K 1,5 mm² (iluminación)", "m"),
    ("el_cable_25", "Cable H07V-K / ES07Z1-K 2,5 mm² (tomas, baños y secadora)", "m"),
    ("el_cable_4", "Cable H07V-K / ES07Z1-K 4 mm² (lavadora, lavavajillas y termo)", "m"),
    ("el_cable_6", "Cable H07V-K / ES07Z1-K 6 mm² (cocina, calefacción y aire acondicionado)", "m"),
    ("el_tubo16", "Tubo corrugado reforzado Ø16 mm", "m"),
    ("el_tubo20", "Tubo corrugado reforzado Ø20 mm", "m"),
    ("el_tubo25", "Tubo corrugado reforzado Ø25 mm", "m"),
    ("el_cuadro", "Cuadro eléctrico CGMP de empotrar con puerta", "ud"),
    ("el_iga_25", "Interruptor general automático (IGA) 25 A", "ud"),
    ("el_iga_40", "Interruptor general automático (IGA) 40 A", "ud"),
    ("el_iga_50", "Interruptor general automático (IGA) 50 A", "ud"),
    ("el_iga_63", "Interruptor general automático (IGA) 63 A", "ud"),
    ("el_dif", "Interruptor diferencial 2P 30 mA tipo A", "ud"),
    ("el_pia_10", "Magnetotérmico (PIA) 10 A", "ud"),
    ("el_pia_16", "Magnetotérmico (PIA) 16 A", "ud"),
    ("el_pia_20", "Magnetotérmico (PIA) 20 A", "ud"),
    ("el_pia_25", "Magnetotérmico (PIA) 25 A", "ud"),
    ("el_spd", "Protector contra sobretensiones (SPD)", "ud"),
    ("el_bornas", "Peines, bornas y repartidores N/PE para el cuadro", "ud"),
    ("el_etiquetas", "Etiquetas de identificación de circuitos", "ud"),
    ("el_interruptor", "Interruptor / conmutador de empotrar", "ud"),
    ("el_base16", "Base de enchufe 16 A 2P+T", "ud"),
    ("el_base25", "Base de enchufe 25 A (cocina / horno)", "ud"),
    ("el_cajas_mec", "Cajas universales de mecanismo", "ud"),
    ("el_marcos", "Marcos y embellecedores", "ud"),
    ("el_cajas_reg", "Cajas de registro / derivación 100×100", "ud"),
    ("el_wago", "Conectores de empalme rápido (2, 3 y 5 vías)", "ud"),
    ("el_terminales", "Terminales y punteras para cable (bolsa)", "bolsa"),
    ("el_cinta", "Cinta aislante", "ud"),
    ("el_bridas", "Bridas de nylon (bolsa de 100)", "bolsa"),
    ("el_abraz", "Abrazaderas / grapas para tubo (bolsa de 100)", "bolsa"),
    ("el_tornillos", "Tornillos de fijación de mecanismos (caja)", "caja"),
    ("el_tacos", "Tacos y tornillos Ø6 mm (caja de 100)", "caja"),
])
_cat(CAP_FONT, [
    ("fo_mc16", "Tubo multicapa PEX-AL-PEX Ø16 mm", "m"),
    ("fo_mc20", "Tubo multicapa PEX-AL-PEX Ø20 mm", "m"),
    ("fo_mc25", "Tubo multicapa PEX-AL-PEX Ø25 mm", "m"),
    ("fo_coquilla", "Coquilla aislante para tubería de agua caliente", "m"),
    ("fo_llave_gen", "Llave de corte general de esfera", "ud"),
    ("fo_llave_local", "Llaves de corte de esfera por local húmedo", "ud"),
    ("fo_escuadra", "Llaves de escuadra 1/2\"-3/8\" con filtro", "ud"),
    ("fo_latiguillo", "Latiguillos flexibles de conexión a aparato", "ud"),
    ("fo_acc_mc", "Accesorios multicapa (codos, tes, manguitos)", "ud"),
    ("fo_teflon", "Cinta de teflón / PTFE", "ud"),
    ("fo_pvc40", "Tubo PVC evacuación Ø40 mm", "m"),
    ("fo_pvc50", "Tubo PVC evacuación Ø50 mm", "m"),
    ("fo_pvc110", "Tubo PVC evacuación Ø110 mm", "m"),
    ("fo_acc_pvc", "Accesorios PVC (codos 45°/90°, tes, manguitos)", "ud"),
    ("fo_sifon_lav", "Sifón de lavabo / bidé", "ud"),
    ("fo_sifon_fre", "Sifón de fregadero", "ud"),
    ("fo_valv_ducha", "Válvula de desagüe de plato de ducha", "ud"),
    ("fo_aireacion", "Válvula de aireación (admisión de aire)", "ud"),
    ("fo_pegamento", "Pegamento para PVC (bote)", "ud"),
    ("fo_abraz_iso", "Abrazaderas isofónicas antivibración (M8)", "ud"),
    ("fo_tacos", "Tacos y tornillos para abrazaderas (caja de 100)", "caja"),
])
_cat(CAP_TELE, [
    ("te_cable", "Cable Cat6 U/UTP LSZH", "m"),
    ("te_tubo", "Tubo corrugado Ø20 mm independiente para datos", "m"),
    ("te_roseta", "Toma / roseta RJ45 Cat6", "ud"),
    ("te_caja", "Caja universal y marco para toma RJ45", "ud"),
    ("te_patch", "Patch panel 24 puertos Cat6 1U", "ud"),
    ("te_rack", "Armario rack 19\" (6U / 9U)", "ud"),
    ("te_switch", "Switch Gigabit (puertos según tomas)", "ud"),
    ("te_latiguillos", "Latiguillos de parcheo Cat6", "ud"),
    ("te_regleta", "Regleta de alimentación 19\" para rack", "ud"),
    ("te_velcro", "Bridas de velcro y organizador de cable", "ud"),
    ("te_etiquetas", "Etiquetas de identificación de tomas", "ud"),
    ("te_bridas", "Bridas, grapas y tornillería de fijación (kit)", "ud"),
])
_cat(CAP_CCTV, [
    ("cc_domo_ip", "Cámara IP domo (interior / techo) PoE", "ud"),
    ("cc_bullet_ip", "Cámara IP bullet (pared / exterior) PoE", "ud"),
    ("cc_ptz_ip", "Cámara IP PTZ motorizada 360° PoE+", "ud"),
    ("cc_domo_hd", "Cámara analógica HD domo (HD-TVI/CVI)", "ud"),
    ("cc_bullet_hd", "Cámara analógica HD bullet (HD-TVI/CVI)", "ud"),
    ("cc_ptz_hd", "Cámara analógica HD PTZ (HD-TVI/CVI)", "ud"),
    ("cc_nvr_4", "Grabador NVR IP 4K de 4 canales", "ud"),
    ("cc_nvr_8", "Grabador NVR IP 4K de 8 canales", "ud"),
    ("cc_nvr_16", "Grabador NVR IP 4K de 16 canales", "ud"),
    ("cc_nvr_32", "Grabador NVR IP 4K de 32 canales", "ud"),
    ("cc_dvr_4", "Grabador DVR HD de 4 canales", "ud"),
    ("cc_dvr_8", "Grabador DVR HD de 8 canales", "ud"),
    ("cc_dvr_16", "Grabador DVR HD de 16 canales", "ud"),
    ("cc_dvr_32", "Grabador DVR HD de 32 canales", "ud"),
    ("cc_hdd_2", "Disco duro Surveillance 24/7 de 2 TB", "ud"),
    ("cc_hdd_4", "Disco duro Surveillance 24/7 de 4 TB", "ud"),
    ("cc_hdd_6", "Disco duro Surveillance 24/7 de 6 TB", "ud"),
    ("cc_hdd_8", "Disco duro Surveillance 24/7 de 8 TB", "ud"),
    ("cc_hdd_12", "Disco duro Surveillance 24/7 de 12 TB", "ud"),
    ("cc_hdd_16", "Disco duro Surveillance 24/7 de 16 TB", "ud"),
    ("cc_poe_4", "Switch PoE+ Gigabit de 4 puertos", "ud"),
    ("cc_poe_8", "Switch PoE+ Gigabit de 8 puertos", "ud"),
    ("cc_poe_16", "Switch PoE+ Gigabit de 16 puertos", "ud"),
    ("cc_poe_24", "Switch PoE+ Gigabit de 24 puertos", "ud"),
    ("cc_fuente", "Fuente de alimentación 12 V multisalida (8 salidas)", "ud"),
    ("cc_cable_ip", "Cable Cat6 UTP para cámaras IP", "m"),
    ("cc_cable_coax", "Cable coaxial RG59 + alimentación 2×0,75 mm² (siamés)", "m"),
    ("cc_tubo", "Tubo corrugado Ø20 mm para CCTV", "m"),
    ("cc_rj45", "Conectores RJ45 Cat6", "ud"),
    ("cc_bnc", "Conectores BNC y conectores de alimentación DC", "ud"),
    ("cc_caja", "Cajas de conexión estancas IP66", "ud"),
    ("cc_soporte", "Soportes de pared / brazos para cámara", "ud"),
    ("cc_tornilleria", "Tornillería, tacos y bridas para cámaras (kit)", "ud"),
    ("cc_monitor", "Monitor LED Full HD 22\"-32\"", "ud"),
    ("cc_hdmi", "Cable HDMI 2.0", "ud"),
    ("cc_cartel", "Cartel informativo de zona videovigilada", "ud"),
])
_cat(CAP_CLIMA, [
    ("cl_equipo_4", "Bomba de calor / aerotermia Inverter 4 kW", "ud"),
    ("cl_equipo_6", "Bomba de calor / aerotermia Inverter 6 kW", "ud"),
    ("cl_equipo_8", "Bomba de calor / aerotermia Inverter 8 kW", "ud"),
    ("cl_equipo_12", "Bomba de calor / aerotermia Inverter 12 kW", "ud"),
    ("cl_equipo_16", "Bomba de calor / aerotermia Inverter 16 kW", "ud"),
    ("cl_termostato", "Termostato de ambiente programable", "ud"),
    ("cl_tubo", "Tubería multicapa Ø32 mm para circuito hidráulico", "m"),
    ("cl_coquilla", "Aislamiento (coquilla) para circuito hidráulico", "m"),
    ("cl_vaso", "Vaso de expansión", "ud"),
    ("cl_valvulas", "Válvulas de corte, purgadores y filtro (kit)", "ud"),
    ("cl_antivib", "Soportes antivibratorios y bancada", "ud"),
])
CLAVES_CATALOGO = {c[0] for c in CATALOGO}
EQUIPOS_CLIMA_KW = [4, 6, 8, 12, 16]


def calcular_materiales(c: dict) -> dict:
    """Cantidades ORIENTATIVAS por clave del catálogo: {clave: (cantidad, nota)}.
    Los recorridos de cable y tubo son estimaciones: ajústalas con tus parámetros de obra."""
    q: dict = {}
    w = 1 + c["desperdicio_pct"] / 100

    def put(clave, cant, nota=""):
        if clave in CLAVES_CATALOGO and cant and cant > 0:
            q[clave] = (math.ceil(cant - 1e-9), nota)

    # ---- Electricidad ----
    if c["inc_elec"] and c["filas_circ"]:
        n_circ = {f["Circuito"]: f["Nº circuitos"] for f in c["filas_circ"]}
        total_circ = sum(n_circ.values())
        luz, c2, c5 = c["mec"]["puntos_luz"], c["mec"]["tomas_c2"], c["mec"]["tomas_c5"]
        mp, md = c["m_punto"], c["m_dedicado"]

        def ded(cod):
            return md if cod in n_circ else 0

        r = {"C1": luz * mp, "C2": c2 * mp, "C5": c5 * mp, "C3": ded("C3"), "C4": ded("C4"),
             "C8": ded("C8"), "C9": ded("C9"), "C10": ded("C10")}
        tres = "3 conductores (fase, neutro, tierra)"
        put("el_cable_15", r["C1"] * 3 * w, tres)
        put("el_cable_25", (r["C2"] + r["C5"] + r["C10"]) * 3 * w, tres)
        put("el_cable_4", r["C4"] * 3 * w, tres)
        put("el_cable_6", (r["C3"] + r["C8"] + r["C9"]) * 3 * w, tres)
        t16 = r["C1"]
        t20 = r["C2"] + r["C4"] + r["C5"] + r["C10"]
        t25 = r["C3"] + r["C8"] + r["C9"]
        put("el_tubo16", t16 * w)
        put("el_tubo20", t20 * w)
        put("el_tubo25", t25 * w)
        tubo_total = t16 + t20 + t25
        if c["di"]:
            put(f"el_di_{int(c['di']['seccion_mm2'])}", c["longitud_di"] * 3 * w,
                "3 conductores (fase, neutro, protección)")
            put("el_tubo_di", c["longitud_di"] * w)
        n_dif = math.ceil(total_circ / 5)
        modulos = math.ceil((2 + 2 * n_dif + total_circ + (2 if c["spd"] else 0)) * 1.2)
        tam = next((m for m in (12, 18, 24, 36, 48) if m >= modulos), 72)
        put("el_cuadro", 1, f"{tam} módulos (necesarios ≈ {modulos}, con 20 % de reserva)")
        iga = next((a for a in IGA_ESTANDAR_A if a >= c["potencia"] / TENSION_V - 1e-6), None)
        if iga:
            put(f"el_iga_{iga}", 1, f"Potencia prevista {c['potencia']:.0f} W")
        put("el_dif", n_dif, "Máximo 5 circuitos por diferencial")
        put("el_pia_10", n_circ.get("C1", 0))
        put("el_pia_16", n_circ.get("C2", 0) + n_circ.get("C5", 0) + n_circ.get("C10", 0))
        put("el_pia_20", n_circ.get("C4", 0))
        put("el_pia_25", n_circ.get("C3", 0) + n_circ.get("C8", 0) + n_circ.get("C9", 0))
        if c["spd"]:
            put("el_spd", 1)
        put("el_bornas", 1)
        put("el_etiquetas", total_circ)
        base16 = c2 + c5 + (3 if "C4" in n_circ else 0) + (1 if "C10" in n_circ else 0)
        base25 = 1 if "C3" in n_circ else 0
        put("el_interruptor", luz)
        put("el_base16", base16)
        put("el_base25", base25)
        cajas = luz + base16 + base25
        put("el_cajas_mec", cajas)
        put("el_marcos", cajas)
        put("el_cajas_reg", total_circ * 2 + c["n_estancias"])
        put("el_wago", (luz + c2 + c5) * 2)
        put("el_terminales", 1)
        put("el_cinta", max(1, c["superficie"] / 40))
        put("el_bridas", max(1, c["superficie"] / 60))
        put("el_abraz", tubo_total / 80)
        put("el_tornillos", max(1, cajas / 50))
        put("el_tacos", max(1, c["superficie"] / 80))

    # ---- Fontanería ----
    if c["inc_font"]:
        ap = c["aparatos"]
        n_ap = sum(ap.values())
        n_loc = c["n_banos"] + c["n_aseos"] + c["n_cocinas"] + c["n_galerias"]
        if n_ap:
            m25, m20, m16 = 4 + 2 * n_loc, 6 * n_loc, 3 * n_ap
            agua = m25 + m20 + m16
            put("fo_mc16", m16 * w, "Ramales a aparatos")
            put("fo_mc20", m20 * w, "Distribución a locales húmedos (fría + caliente)")
            put("fo_mc25", m25 * w, "Alimentación general")
            put("fo_coquilla", (0.5 * m20 + 0.4 * m16) * w, "Estimación para agua caliente")
            put("fo_llave_gen", 1)
            put("fo_llave_local", 2 * n_loc)
            esc = n_ap - ap["ducha"]
            put("fo_escuadra", esc)
            put("fo_latiguillo", esc + ap["lavabo"] + ap["fregadero"])
            put("fo_acc_mc", agua / 2)
            put("fo_teflon", max(1, n_ap / 6))
            pvc40, pvc50 = 1.5 * ap["lavabo"], 2.5 * (ap["ducha"] + ap["fregadero"] + ap["lavadora"])
            pvc110 = 3 + 2 * ap["inodoro"]
            put("fo_pvc40", pvc40 * w)
            put("fo_pvc50", pvc50 * w)
            put("fo_pvc110", pvc110 * w)
            pvc = pvc40 + pvc50 + pvc110
            put("fo_acc_pvc", pvc / 1.5)
            put("fo_sifon_lav", ap["lavabo"])
            put("fo_sifon_fre", ap["fregadero"])
            put("fo_valv_ducha", ap["ducha"])
            put("fo_aireacion", c["n_banos"] + c["n_aseos"] + c["n_cocinas"])
            put("fo_pegamento", max(1, pvc / 20))
            abraz = math.ceil(pvc / 0.8) + math.ceil(agua / 1.2)
            put("fo_abraz_iso", abraz)
            put("fo_tacos", max(1, abraz / 100))

    # ---- Telecomunicaciones ----
    tomas = c["tomas_datos"]
    if c["inc_teleco"] and tomas > 0:
        m = tomas * c["m_toma_datos"]
        put("te_cable", m * w)
        put("te_tubo", m)
        put("te_roseta", tomas)
        put("te_caja", tomas)
        put("te_patch", 1)
        put("te_rack", 1)
        puertos = next((p for p in (8, 16, 24, 48) if p >= tomas), 48)
        put("te_switch", 1, f"{puertos} puertos")
        put("te_latiguillos", tomas)
        put("te_regleta", 1)
        put("te_velcro", 1)
        put("te_etiquetas", tomas)
        put("te_bridas", 1)

    # ---- CCTV ----
    cc = c["cctv"]
    if c["inc_cctv"] and cc and cc["n_cam"] > 0:
        suf = "ip" if cc["ip"] else "hd"
        put(f"cc_domo_{suf}", cc["n_domo"])
        put(f"cc_bullet_{suf}", cc["n_bullet"])
        put(f"cc_ptz_{suf}", cc["n_ptz"])
        put(f"cc_{'nvr' if cc['ip'] else 'dvr'}_{cc['canales']}", cc["n_grabadores"])
        put(f"cc_hdd_{cc['disco_tb']}", cc["n_discos"],
            f"≈ {cc['tb']:.1f} TB para {cc['dias']} días")
        n_cam = cc["n_cam"]
        if cc["ip"]:
            put(f"cc_poe_{cc['puertos_poe']}", cc["n_switch"],
                f"Consumo PoE estimado ≈ {cc['poe_total_w']:.0f} W")
            put("cc_cable_ip", cc["cable_m"] * w)
            put("cc_rj45", n_cam * 2)
        else:
            put("cc_fuente", cc["fuentes"])
            put("cc_cable_coax", cc["cable_m"] * w)
            put("cc_bnc", n_cam * 3)
        put("cc_tubo", cc["cable_m"])
        put("cc_caja", n_cam)
        put("cc_soporte", cc["n_bullet"] + cc["n_ptz"])
        put("cc_tornilleria", n_cam)
        if cc["monitor"]:
            put("cc_monitor", 1)
            put("cc_hdmi", 1)
        put("cc_cartel", 1)

    # ---- Climatización ----
    cl = c["cl"]
    if c["inc_clima"] and cl and c["superficie"] > 0:
        kw = cl["kw"]
        clase = next((k for k in EQUIPOS_CLIMA_KW if k >= kw), EQUIPOS_CLIMA_KW[-1])
        n_eq = 1 if kw <= EQUIPOS_CLIMA_KW[-1] else math.ceil(kw / EQUIPOS_CLIMA_KW[-1])
        put(f"cl_equipo_{clase}", n_eq, f"Potencia estimada {kw:.1f} kW")
        put("cl_termostato", max(1, c["n_dormitorios"] + c["n_salones"]))
        m = 10 + 0.2 * c["superficie"]
        put("cl_tubo", m * w)
        put("cl_coquilla", m * w)
        put("cl_vaso", 1)
        put("cl_valvulas", 1)
        put("cl_antivib", 1)
    return q


def listado_materiales(cantidades: dict, precios: dict) -> list[dict]:
    """Filas del listado de materiales con precios (0 € si no se han indicado)."""
    filas = []
    for clave, cap, nombre, ud in CATALOGO:
        if clave in cantidades:
            cant, nota = cantidades[clave]
            pu = float(precios.get(clave, 0.0) or 0.0)
            filas.append({"Capítulo": cap, "Material": nombre, "Ud": ud, "Cantidad": cant,
                          "Precio unit. (€)": pu, "Importe (€)": cant * pu, "Nota": nota})
    return filas


# ----------------------------------------------------------------------------
# PRESUPUESTO
# ----------------------------------------------------------------------------
def calc_w_m2(d) -> float:
    return RATIOS_W_M2[d["aislamiento_termico"]]


def partidas(d) -> list[dict]:
    """Partidas del presupuesto por tarifas (€/m², €/ud) a partir de un mapeo."""
    sup = d["superficie"]
    out = []

    def add(cap, texto, ud, cant, precio):
        out.append({"Capítulo": cap, "Partida": texto, "Ud": ud, "Cantidad": cant,
                    "Precio unit. (€)": precio, "Importe (€)": cant * precio})

    if d["inc_elec"]:
        add(CAP_ELEC, "Instalación eléctrica general", "m²", sup, d["pu_elec_m2"])
        add(CAP_ELEC, "Cuadro general de mando y protección", "ud",
            1 if sup > 0 else 0, d["pu_elec_cgmp"])
        if d["incluir_spd"]:
            add(CAP_ELEC, "Protector contra sobretensiones", "ud", 1, d["pu_elec_spd"])
    if d["inc_font"]:
        add(CAP_FONT, "Instalación de baño completo", "ud", d["n_banos"], d["pu_font_bano"])
        add(CAP_FONT, "Instalación de aseo", "ud", d["n_aseos"], d["pu_font_aseo"])
        add(CAP_FONT, "Instalación de cocina", "ud", d["n_cocinas"], d["pu_font_cocina"])
    if d["inc_teleco"]:
        tomas = d["n_dormitorios"] * d["tomas_datos_dorm"] + d["n_salones"] * d["tomas_datos_salon"]
        add(CAP_TELE, "Toma de datos RJ45 Cat6", "toma", tomas, d["pu_teleco_toma"])
    if d["inc_cctv"]:
        fijas = d["cctv_domo"] + d["cctv_bullet"]
        hay = 1 if fijas + d["cctv_ptz"] > 0 else 0
        add(CAP_CCTV, "Cámara fija (domo / bullet) instalada", "ud", fijas, d["pu_cctv_camara"])
        add(CAP_CCTV, "Cámara PTZ instalada", "ud", d["cctv_ptz"], d["pu_cctv_ptz"])
        add(CAP_CCTV, "Grabador NVR/DVR con disco duro", "ud", hay, d["pu_cctv_grabador"])
        add(CAP_CCTV, "Monitor de visualización", "ud", hay if d["cctv_monitor"] else 0,
            d["pu_cctv_monitor"])
    if d["inc_clima"]:
        kw = climatizacion(sup, calc_w_m2(d))["kw"]
        add(CAP_CLIMA, "Instalación de climatización", "kW", kw, d["pu_clima_kw"])
    return out


def armar_presupuesto(lineas_tarifas: list[dict], materiales: list[dict], metodo: str,
                      mo_pct: float, iva_pct: float, incluidos: dict) -> dict:
    """Resumen por capítulo y totales. metodo: 'tarifas' o 'materiales'."""
    caps = []
    for cap, clave, nombre in CAPITULOS:
        if not incluidos.get(clave):
            continue
        if metodo == "materiales":
            mat = sum(f["Importe (€)"] for f in materiales if f["Capítulo"] == cap)
            mo = mat * mo_pct / 100
            caps.append({"cap": cap, "nombre": nombre, "materiales": mat,
                         "mano_obra": mo, "total": mat + mo})
        else:
            tot = sum(f["Importe (€)"] for f in lineas_tarifas if f["Capítulo"] == cap)
            caps.append({"cap": cap, "nombre": nombre, "materiales": None,
                         "mano_obra": None, "total": tot})
    base = sum(c["total"] for c in caps)
    iva = base * iva_pct / 100
    res = {"metodo": metodo, "mo_pct": mo_pct, "caps": caps, "base": base, "iva_pct": iva_pct,
           "iva": iva, "total": base + iva,
           "lineas": lineas_tarifas if metodo == "tarifas" else materiales}
    if metodo == "materiales":
        res["materiales_total"] = sum(c["materiales"] for c in caps)
        res["mano_obra_total"] = sum(c["mano_obra"] for c in caps)
    return res


# ----------------------------------------------------------------------------
# BIBLIOTECA: normas y manuales (enlaces oficiales; copias locales opcionales)
# ----------------------------------------------------------------------------
def _boe(anio, ident):
    return (f"https://www.boe.es/buscar/act.php?id={ident}",
            f"https://www.boe.es/buscar/pdf/{anio}/{ident}-consolidado.pdf")


_REBT = _boe(2002, "BOE-A-2002-18099")
_CTE = _boe(2006, "BOE-A-2006-5515")
_RITE = _boe(2007, "BOE-A-2007-15820")
_ICT = _boe(2011, "BOE-A-2011-5037")
_SEG = _boe(2014, "BOE-A-2014-3649")
_LOPD = _boe(2018, "BOE-A-2018-16673")
IND = "https://industria.gob.es/Calidad-Industrial/seguridadindustrial/instalacionesindustriales/baja-tension/Paginas/"

BIBLIOTECA = [
    {"id": "rebt", "ambito": "Electricidad", "tipo": "Norma",
     "titulo": "REBT: Reglamento Electrotécnico para Baja Tensión (RD 842/2002)",
     "desc": "Reglamento e instrucciones técnicas complementarias (ITC-BT).",
     "ver": _REBT[0], "pdf": _REBT[1], "archivo": "REBT_RD842-2002.pdf"},
    {"id": "guia_rebt", "ambito": "Electricidad", "tipo": "Guía",
     "titulo": "Guías técnicas de aplicación del REBT (Ministerio de Industria)",
     "desc": "Guías no vinculantes para aplicar el reglamento y sus ITC.",
     "ver": IND + "guia-tecnica-aplicacion.aspx", "pdf": None, "archivo": ""},
    {"id": "faq_rebt", "ambito": "Electricidad", "tipo": "Guía",
     "titulo": "Preguntas frecuentes del REBT (Ministerio de Industria)",
     "desc": "Respuestas oficiales a dudas habituales de instaladores.",
     "ver": IND + "preguntas-frecuentes.aspx", "pdf": None, "archivo": ""},
    {"id": "cte", "ambito": "Fontanería", "tipo": "Norma",
     "titulo": "CTE: Código Técnico de la Edificación (RD 314/2006)",
     "desc": "Parte I del CTE; los documentos básicos se consultan en codigotecnico.org.",
     "ver": _CTE[0], "pdf": _CTE[1], "archivo": "CTE_RD314-2006.pdf"},
    {"id": "dbhs", "ambito": "Fontanería", "tipo": "Norma",
     "titulo": "DB-HS Salubridad (HS4 suministro de agua, HS5 evacuación de aguas)",
     "desc": "Documento básico oficial, con versión comentada del Ministerio.",
     "ver": "https://www.codigotecnico.org/DocumentosCTE/Salubridad.html", "pdf": None,
     "archivo": "DB-HS_Salubridad.pdf"},
    {"id": "rite", "ambito": "Climatización", "tipo": "Norma",
     "titulo": "RITE: Reglamento de Instalaciones Térmicas en los Edificios (RD 1027/2007)",
     "desc": "Texto consolidado del reglamento.",
     "ver": _RITE[0], "pdf": _RITE[1], "archivo": "RITE_RD1027-2007.pdf"},
    {"id": "rite_2021", "ambito": "Climatización", "tipo": "Norma",
     "titulo": "RD 178/2021, por el que se modifica el RITE",
     "desc": "Modificación del RITE vigente desde el 1 de julio de 2021.",
     "ver": "https://www.boe.es/diario_boe/txt.php?id=BOE-A-2021-4572",
     "pdf": "https://www.boe.es/boe/dias/2021/03/24/pdfs/BOE-A-2021-4572.pdf",
     "archivo": "RITE_RD178-2021.pdf"},
    {"id": "ict", "ambito": "Telecomunicaciones", "tipo": "Norma",
     "titulo": "Reglamento ICT (RD 346/2011)",
     "desc": "Infraestructuras comunes de telecomunicaciones en edificios.",
     "ver": _ICT[0], "pdf": _ICT[1], "archivo": "ICT_RD346-2011.pdf"},
    {"id": "aepd_guia", "ambito": "CCTV", "tipo": "Guía",
     "titulo": "AEPD: Guía sobre el uso de videocámaras para seguridad y otras finalidades",
     "desc": "Obligaciones de protección de datos al instalar videovigilancia.",
     "ver": "https://www.aepd.es/areas/videovigilancia/index.html",
     "pdf": "https://www.aepd.es/guias/guia-videovigilancia.pdf",
     "archivo": "AEPD_Guia_videovigilancia.pdf"},
    {"id": "seg_privada", "ambito": "CCTV", "tipo": "Norma",
     "titulo": "Ley 5/2014 de Seguridad Privada",
     "desc": "Regula, entre otras cosas, la instalación de sistemas conectados a centrales de alarma.",
     "ver": _SEG[0], "pdf": _SEG[1], "archivo": "Ley5-2014_Seguridad_Privada.pdf"},
    {"id": "lopdgdd", "ambito": "CCTV", "tipo": "Norma",
     "titulo": "Ley Orgánica 3/2018 de Protección de Datos (LOPDGDD)",
     "desc": "Incluye el tratamiento con fines de videovigilancia.",
     "ver": _LOPD[0], "pdf": _LOPD[1], "archivo": "LO3-2018_LOPDGDD.pdf"},
]
CARPETA_BIBLIOTECA = Path(__file__).resolve().parent / "biblioteca"


# ----------------------------------------------------------------------------
# DOCUMENTOS: contenido común (bloques) + generadores Word, PDF y Excel
# ----------------------------------------------------------------------------
AZUL, AMARILLO, GRIS, CLARO = "0B2545", "FFC400", "5B7088", "F3F8FC"


def fnum(v, dec: int = 2) -> str:
    """Número con formato español: 1.234,56"""
    s = f"{v:,.{dec}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def feur(v) -> str:
    return fnum(v) + " €"


def _tabla(cols, rows, w=None, num=(), resaltar=0):
    return ("tabla", {"cols": cols, "rows": rows, "w": w, "num": tuple(num), "resaltar": resaltar})


def _tabla_materiales(d, cap, filas, con_precios):
    """Bloques con la tabla de materiales de un capítulo."""
    if con_precios:
        cols = ["Material", "Ud", "Cantidad", "Precio unit.", "Importe", "Observaciones"]
        rows = [[f["Material"], f["Ud"], fnum(f["Cantidad"], 0), feur(f["Precio unit. (€)"]),
                 feur(f["Importe (€)"]), f["Nota"]] for f in filas]
        rows.append([f"Subtotal materiales: {NOMBRE_CAP[cap]}", "", "", "",
                     feur(sum(f["Importe (€)"] for f in filas)), ""])
        return _tabla(cols, rows, w=[5.6, 1.0, 1.8, 1.9, 2.0, 4.7], num=(2, 3, 4), resaltar=1)
    cols = ["Material", "Ud", "Cantidad", "Observaciones"]
    rows = [[f["Material"], f["Ud"], fnum(f["Cantidad"], 0), f["Nota"]] for f in filas]
    return _tabla(cols, rows, w=[8.2, 1.2, 1.8, 5.8], num=(2,))


def bloques_listado(d, con_titulo_h1=True):
    """Listado de materiales por capítulo."""
    B = []
    mats = d["mats"]
    con_precios = d["precios_en_uso"]
    for cap, clave, nombre in CAPITULOS:
        filas = [f for f in mats if f["Capítulo"] == cap]
        if not filas:
            continue
        B.append(("h2", nombre))
        B.append(_tabla_materiales(d, cap, filas, con_precios))
    if not B:
        B.append(("p", "No hay materiales calculados: añade estancias y especialidades."))
    B.append(("nota", "Cantidades orientativas calculadas con los parámetros de obra indicados "
                      "(recorridos medios, desperdicio y reglas de estimación). Revísalas antes "
                      "de comprar."))
    return B


def _tabla_presupuesto(d):
    pres = d["pres"]
    iva_txt = f"IVA ({fnum(pres['iva_pct'], 0)} %)"
    if pres["metodo"] == "materiales":
        rows = [[f"Capítulo {i}: {c['nombre']}", feur(c["materiales"]), feur(c["mano_obra"]),
                 feur(c["total"])] for i, c in enumerate(pres["caps"], start=1)]
        rows += [["BASE IMPONIBLE (materiales + mano de obra)", feur(pres["materiales_total"]),
                  feur(pres["mano_obra_total"]), feur(pres["base"])],
                 [iva_txt, "", "", feur(pres["iva"])],
                 ["PRESUPUESTO TOTAL (IVA incluido)", "", "", feur(pres["total"])]]
        return _tabla(["Capítulo / módulo", "Materiales", "Mano de obra", "Total capítulo"], rows,
                      w=[8.0, 3.0, 3.0, 3.0], num=(1, 2, 3), resaltar=1)
    rows = [[f["Capítulo"], f["Partida"], f["Ud"], fnum(f["Cantidad"]), feur(f["Precio unit. (€)"]),
             feur(f["Importe (€)"])] for f in pres["lineas"]]
    rows += [["", "BASE IMPONIBLE", "", "", "", feur(pres["base"])],
             ["", iva_txt, "", "", "", feur(pres["iva"])],
             ["", "PRESUPUESTO TOTAL (IVA incluido)", "", "", "", feur(pres["total"])]]
    return _tabla(["Capítulo", "Partida", "Ud", "Cantidad", "Precio unit.", "Importe"], rows,
                  w=[2.6, 6.2, 1.0, 2.0, 2.5, 2.7], num=(3, 4, 5), resaltar=1)


def _nota_mo(d):
    pres = d["pres"]
    if pres["metodo"] == "materiales":
        return (f"Mano de obra calculada como el {fnum(pres['mo_pct'], 0)} % sobre el importe de "
                "materiales.")
    return "Presupuesto elaborado por partidas con tarifas propias (materiales y mano de obra incluidos)."


def bloques_presupuesto(d):
    B = [("portada", {
        "titulo": "PRESUPUESTO DE REFORMA",
        "subtitulo": "Resumen económico por capítulos",
        "lineas": [f"Cliente: {d['cliente'] or 'Sin especificar'}",
                   f"Obra: {d['obra'] or 'Sin especificar'}",
                   f"Fecha: {d['fecha']:%d/%m/%Y}   ·   Ref.: {d['ref']}"]})]
    B.append(("h1", "Estado de mediciones y presupuesto económico"))
    B.append(_tabla_presupuesto(d))
    B.append(("nota", _nota_mo(d)))
    B.append(("p", f"El presente presupuesto se emite con una validez de {d['validez_dias']} días "
                   "naturales. Las instalaciones se ejecutarán conforme a la normativa vigente y "
                   "se entregarán con la documentación y los certificados que correspondan."))
    return B


def bloques_memoria(d):
    """Memoria técnica y presupuesto: todo el proyecto en un documento."""
    inc = d["inc"]
    B = []
    normas = []
    if inc["inc_elec"]:
        normas.append("REBT (RD 842/2002)")
    if inc["inc_font"]:
        normas.append("CTE DB-HS4/HS5")
    if inc["inc_teleco"]:
        normas.append("ICT (RD 346/2011)")
    if inc["inc_clima"]:
        normas.append("RITE (RD 1027/2007, mod. RD 178/2021)")
    if inc["inc_cctv"]:
        normas.append("RGPD / LOPDGDD")
    lineas = []
    if normas:
        lineas.append("Normativa de aplicación: " + " | ".join(normas))
    lineas += [f"Cliente: {d['cliente'] or 'Sin especificar'}",
               f"Obra: {d['obra'] or 'Sin especificar'}",
               f"Fecha: {d['fecha']:%d/%m/%Y}   ·   Ref.: {d['ref']}"]
    B.append(("portada", {
        "titulo": "PROYECTO TÉCNICO DE REFORMAS E INSTALACIONES",
        "subtitulo": "Memoria técnica justificativa, especificaciones de cálculo y estado de mediciones",
        "lineas": lineas}))

    n = 1
    # 1. Resumen y alcance
    B.append(("h1", f"{n}. Resumen ejecutivo y alcance del proyecto"))
    esp = [nombre for cap, clave, nombre in CAPITULOS if inc[clave]]
    B.append(("p", "El presente documento constituye la memoria técnica y justificativa del proyecto "
                   "de reformas e instalaciones descrito a continuación. Los resultados de cálculo "
                   "son orientativos y deben ser validados por el instalador autorizado o técnico "
                   "competente responsable de la obra."))
    datos = [["Cliente", d["cliente"] or "Sin especificar"],
             ["Dirección de la obra", d["obra"] or "Sin especificar"],
             ["Superficie útil afectada", f"{fnum(d['superficie'], 1)} m²"],
             ["Especialidades incluidas", "; ".join(esp) or "Ninguna"]]
    B.append(_tabla(["Dato", "Valor"], datos, w=[5, 12]))
    est = [[etq, str(num)] for etq, num in d["estancias"] if num]
    if est:
        B.append(("h2", "Estancias consideradas"))
        B.append(_tabla(["Estancia", "Cantidad"], est, w=[12, 5], num=(1,)))

    # Electricidad
    if inc["inc_elec"] and d["elec"]:
        e = d["elec"]
        n += 1
        B.append(("h1", f"{n}. Instalación eléctrica y cuadro general (REBT)"))
        B.append(("p", "Conforme a las instrucciones técnicas complementarias ITC-BT-15, ITC-BT-17, "
                       "ITC-BT-19, ITC-BT-23 e ITC-BT-25 del Reglamento Electrotécnico para Baja "
                       "Tensión (REBT, RD 842/2002):"))
        pts = [f"Electrificación {e['grado'].lower()} con potencia prevista de "
               f"{fnum(e['potencia'] / 1000, 2)} kW (IGA de {e['iga']} A)."]
        di = e["di"]
        if di:
            pts.append(
                f"Derivación individual: conductores de cobre unipolares de {di['seccion_mm2']:g} mm² "
                f"(libres de halógenos, ES07Z1-K) alojados en tubo, con una longitud de "
                f"{fnum(e['longitud_di'], 1)} m, intensidad de {fnum(di['intensidad_a'], 1)} A y caída "
                f"de tensión de {fnum(di['caida_pct'])} % (límite {fnum(e['limite'], 1)} %: "
                f"{e['suministro'].split(' (')[0].lower()}).")
        pts.append("Cuadro general de mando y protección (CGMP) dimensionado con los circuitos de la "
                   "tabla siguiente, con la regla de máximo 5 circuitos por interruptor diferencial "
                   "de 30 mA." + (" Incluye protector contra sobretensiones." if e["spd"] else ""))
        B.append(("bullets", pts))
        if e["filas_circ"]:
            B.append(("h2", "Circuitos y puntos mínimos (ITC-BT-25)"))
            filas = [[f["Circuito"], f["Descripción"], str(f["Nº circuitos"]), str(f["PIA (A)"]),
                      f"{f['Sección (mm²)']:g}".replace(".", ","), str(f["Tubo (mm)"]), str(f["Puntos previstos"])]
                     for f in e["filas_circ"]]
            B.append(_tabla(["Circ.", "Descripción", "Nº", "PIA (A)", "Sección (mm²)", "Tubo (mm)",
                             "Puntos"], filas, w=[1.2, 6.0, 1.0, 1.7, 2.3, 2.0, 1.8],
                            num=(2, 3, 4, 5, 6)))
            B.append(("h2", "Protecciones del cuadro"))
            B.append(_tabla(["Elemento", "Valor", "Observaciones"],
                            [[p["Elemento"], p["Valor"], p["Observaciones"]] for p in e["prot"]],
                            w=[4.5, 2.5, 10]))
        B.append(("nota", "Valores orientativos; los mínimos por estancia son una versión simplificada "
                          "de la tabla 2 de la ITC-BT-25."))

    # Fontanería
    if inc["inc_font"] and d["font"]:
        h = d["font"]["hid"]
        n += 1
        B.append(("h1", f"{n}. Fontanería y evacuación de aguas (CTE DB-HS4 / HS5)"))
        B.append(("p", "Seguridad higiénica y evacuación eficiente según el Código Técnico de la "
                       "Edificación:"))
        pts = []
        if h["n_aparatos"]:
            pts.append(f"Red de agua potable (DB-HS4): {h['n_aparatos']} aparatos; caudal instalado de "
                       f"{fnum(h['q_fria'])} L/s en agua fría, coeficiente de simultaneidad "
                       f"{fnum(h['kn'])} (K = 1/raíz(N-1)) y caudal de diseño de {fnum(h['q_diseno'])} "
                       f"L/s. Diámetro interior mínimo de la tubería general: "
                       f"{fnum(h['diametro_mm'], 1)} mm (velocidad de cálculo "
                       f"{fnum(d['font']['velocidad'], 1)} m/s).")
        pts.append("Distribución mediante tubería multicapa o cobre, con llaves de escuadra con "
                   "filtro en cada aparato y llaves de corte general y por local húmedo.")
        pts.append("Saneamiento y evacuación (DB-HS5): colectores en PVC con pendiente adecuada, "
                   "abrazaderas isofónicas, codos de paso suave y sifones con válvulas de aireación.")
        pts.append("Presión en puntos de consumo: mínimo 100 kPa en grifos y 150 kPa en fluxores y "
                   "calentadores; máximo 500 kPa.")
        B.append(("bullets", pts))

    # Telecomunicaciones
    if inc["inc_teleco"] and d["tele"]:
        t = d["tele"]
        n += 1
        B.append(("h1", f"{n}. Telecomunicaciones y red de datos"))
        B.append(("p", f"Cableado estructurado de categoría 6 con {t['tomas']} tomas de datos "
                       f"({t['td_dorm']} por dormitorio y {t['td_salon']} por salón), canalizado en "
                       "tubo corrugado de Ø20 mm independiente de la instalación eléctrica y "
                       "centralizado en armario rack de 19\" con patch panel y switch Gigabit."))

    # CCTV
    if inc["inc_cctv"] and d["cctv"]:
        cc = d["cctv"]
        n += 1
        B.append(("h1", f"{n}. Videovigilancia y CCTV"))
        if cc["n_cam"] <= 0:
            B.append(("p", "Sistema de videovigilancia previsto sin cámaras definidas todavía."))
        else:
            B.append(("p", f"Sistema de seguridad en tecnología {cc['tecnologia']} con "
                           f"{cc['n_cam']} cámaras: {cc['n_domo']} domo (techo/interior), "
                           f"{cc['n_bullet']} bullet (pared/exterior) y {cc['n_ptz']} PTZ "
                           "(motorizadas 360°)."))
            pts = [f"Grabador {'NVR' if cc['ip'] else 'DVR'} de {cc['canales']} canales"
                   + (f" ({cc['n_grabadores']} unidades)." if cc["n_grabadores"] > 1 else "."),
                   f"Almacenamiento: {fnum(cc['gb_dia'], 0)} GB/día con {fnum(cc['mbps'], 0)} Mbps por "
                   f"cámara; {cc['dias']} días de grabación requieren ≈ {fnum(cc['tb'], 1)} TB "
                   f"({cc['n_discos']} × disco Surveillance de {cc['disco_tb']} TB).",
                   f"Cableado: {fnum(cc['cable_m'], 0)} m ({fnum(cc['dist_m'], 0)} m de media por "
                   f"cámara) en tubo corrugado Ø20 mm."]
            if cc["ip"]:
                pts.append(f"Alimentación PoE: consumo estimado de {fnum(cc['poe_total_w'], 0)} W "
                           f"(con margen del 25 %); switch PoE+ de {cc['puertos_poe']} puertos.")
            else:
                pts.append(f"Alimentación 12 V centralizada: {cc['fuentes']} fuente(s) multisalida.")
            if cc["monitor"]:
                pts.append("Visualización local en monitor LED Full HD.")
            B.append(("bullets", pts))
        B.append(("nota", "Instala el cartel informativo de zona videovigilada y cumple el RGPD y la "
                          "LOPDGDD. Si el sistema se conecta a una central receptora de alarmas, la "
                          "instalación debe hacerla una empresa de seguridad privada habilitada "
                          "(Ley 5/2014)."))

    # Climatización
    if inc["inc_clima"] and d["clima"]:
        cl = d["clima"]
        n += 1
        B.append(("h1", f"{n}. Climatización (RITE)"))
        B.append(("p", f"Potencia térmica estimada de {fnum(cl['cl']['kw'])} kW "
                       f"({fnum(cl['cl']['kcal_h'], 0)} kcal/h) con aislamiento «{cl['ratio']}» sobre "
                       f"{fnum(d['superficie'], 1)} m². Es una estimación por ratio: el cálculo de "
                       "cargas térmicas según el RITE debe realizarlo el técnico competente."))

    # Presupuesto
    n += 1
    B.append(("h1", f"{n}. Estado de mediciones y presupuesto económico"))
    B.append(("p", "A continuación se presenta el resumen económico consolidado de las mediciones "
                   "de la obra:"))
    B.append(_tabla_presupuesto(d))
    B.append(("nota", _nota_mo(d)))
    B.append(("p", f"El presente presupuesto se emite con una validez de {d['validez_dias']} días "
                   "naturales. Todas las instalaciones se ejecutarán conforme a la normativa vigente "
                   "y se entregarán con la documentación y los certificados que correspondan "
                   "(por ejemplo, el certificado de instalación eléctrica y los boletines)."))

    # Anexo materiales
    if d["mats"]:
        n += 1
        B.append(("h1", f"{n}. Anexo: listado de materiales"))
        B += bloques_listado(d)

    # Normativa
    n += 1
    B.append(("h1", f"{n}. Normativa de referencia"))
    refs = [f"{b['titulo']}" for b in BIBLIOTECA
            if (b["ambito"] == "Electricidad" and inc["inc_elec"])
            or (b["ambito"] == "Fontanería" and inc["inc_font"])
            or (b["ambito"] == "Climatización" and inc["inc_clima"])
            or (b["ambito"] == "Telecomunicaciones" and inc["inc_teleco"])
            or (b["ambito"] == "CCTV" and inc["inc_cctv"])]
    if refs:
        B.append(("bullets", refs))
    B.append(("nota", f"Estado de la normativa a {NORMATIVA_FECHA}. Comprueba el BOE antes de cada "
                      "proyecto: hay modificaciones del REBT, del RITE y del CTE en revisión."))
    return B


# ---------------------------------------------------------------- Word (.docx)
def _fuente_docx(estilo, nombre="Arial"):
    estilo.font.name = nombre
    rpr = estilo.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), nombre)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]


def _sombrear(celda, color):
    tc_pr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def _borde_inferior(par, color=AMARILLO, sz=12):
    p_pr = par._p.get_or_add_pPr()
    pb = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), str(sz))
    b.set(qn("w:space"), "1")
    b.set(qn("w:color"), color)
    pb.append(b)
    p_pr.append(pb)


def _campo_docx(par, instr, tam=8):
    def fc(tipo):
        r = par.add_run()
        r.font.size = Pt(tam)
        e = OxmlElement("w:fldChar")
        e.set(qn("w:fldCharType"), tipo)
        r._r.append(e)

    fc("begin")
    r = par.add_run()
    r.font.size = Pt(tam)
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    r._r.append(it)
    fc("separate")
    r = par.add_run("1")
    r.font.size = Pt(tam)
    fc("end")


def _tabla_docx(doc, t, ancho):
    cols, rows = t["cols"], t["rows"]
    pesos = t["w"] or [1] * len(cols)
    tot = sum(pesos)
    anchos = [ancho * p / tot for p in pesos]
    num = set(t["num"])
    tabla = doc.add_table(rows=1, cols=len(cols))
    tabla.style = "Table Grid"
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabla.autofit = False
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tabla._tbl.tblPr.append(lay)
    for i, ancho_col in enumerate(anchos):
        tabla.columns[i].width = Cm(ancho_col)
    fila0 = tabla.rows[0]
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    fila0._tr.get_or_add_trPr().append(th)
    for i, c in enumerate(cols):
        celda = fila0.cells[i]
        celda.width = Cm(anchos[i])
        _sombrear(celda, AZUL)
        p = celda.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if i in num else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(c)
        r.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(255, 255, 255)
    n_res = t["resaltar"]
    for ri, fila in enumerate(rows):
        cells = tabla.add_row().cells
        no_cortar = OxmlElement("w:cantSplit")
        tabla.rows[-1]._tr.get_or_add_trPr().append(no_cortar)
        resaltada = n_res and ri >= len(rows) - n_res
        for i, v in enumerate(fila):
            celda = cells[i]
            celda.width = Cm(anchos[i])
            if resaltada:
                _sombrear(celda, "FFE9A6")
            elif ri % 2 == 1:
                _sombrear(celda, CLARO)
            p = celda.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if i in num else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(v))
            r.font.size = Pt(8.5)
            r.bold = bool(resaltada)
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(2)


def render_docx(bloques, d, titulo_doc) -> bytes:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin, sec.bottom_margin = Cm(2.8), Cm(2.0)
    sec.header_distance, sec.footer_distance = Cm(1.0), Cm(0.9)
    ancho = 17.0
    normal = doc.styles["Normal"]
    _fuente_docx(normal)
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(4)
    for nombre, tam in (("Heading 1", 13), ("Heading 2", 11)):
        e = doc.styles[nombre]
        _fuente_docx(e)
        e.font.size = Pt(tam)
        e.font.bold = True
        e.font.color.rgb = RGBColor.from_string(AZUL)
        e.paragraph_format.space_before = Pt(12 if tam == 13 else 8)
        e.paragraph_format.space_after = Pt(4)
        e.paragraph_format.keep_with_next = True
    _fuente_docx(doc.styles["List Bullet"])

    emp = d["empresa"]
    hp = sec.header.paragraphs[0]
    hp.style = doc.styles["Normal"]
    hp.paragraph_format.tab_stops.add_tab_stop(Cm(ancho), WD_TAB_ALIGNMENT.RIGHT)
    r = hp.add_run(emp["nombre"] or "PRO-REFORMAS")
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor.from_string(AZUL)
    r = hp.add_run("\t" + titulo_doc)
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string(GRIS)
    h2 = sec.header.add_paragraph(style=doc.styles["Normal"])
    h2.paragraph_format.tab_stops.add_tab_stop(Cm(ancho), WD_TAB_ALIGNMENT.RIGHT)
    r = h2.add_run(" · ".join(x for x in (emp["nif"], emp["contacto"]) if x))
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(GRIS)
    r = h2.add_run(f"\tRef. {d['ref']} · {d['fecha']:%d/%m/%Y}")
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(GRIS)
    _borde_inferior(h2)

    fp = sec.footer.paragraphs[0]
    fp.style = doc.styles["Normal"]
    fp.paragraph_format.tab_stops.add_tab_stop(Cm(ancho), WD_TAB_ALIGNMENT.RIGHT)
    r = fp.add_run("Documento orientativo: no sustituye el proyecto ni la memoria técnica de un "
                   "técnico competente.")
    r.font.size = Pt(7.5)
    r.font.color.rgb = RGBColor.from_string(GRIS)
    r = fp.add_run("\tPágina ")
    r.font.size = Pt(8)
    _campo_docx(fp, "PAGE")
    r = fp.add_run(" de ")
    r.font.size = Pt(8)
    _campo_docx(fp, "NUMPAGES")

    for tipo, val in bloques:
        if tipo == "portada":
            p = doc.add_paragraph()
            r = p.add_run(val["titulo"])
            r.bold = True
            r.font.size = Pt(18)
            r.font.color.rgb = RGBColor.from_string(AZUL)
            p = doc.add_paragraph()
            r = p.add_run(val["subtitulo"])
            r.italic = True
            r.font.size = Pt(11)
            r.font.color.rgb = RGBColor.from_string(GRIS)
            for linea in val["lineas"]:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(1)
                r = p.add_run(linea)
                r.font.size = Pt(9.5)
            _borde_inferior(p)
        elif tipo == "h1":
            doc.add_heading(val, level=1)
        elif tipo == "h2":
            doc.add_heading(val, level=2)
        elif tipo == "p":
            doc.add_paragraph(val)
        elif tipo == "bullets":
            for item in val:
                doc.add_paragraph(item, style="List Bullet")
        elif tipo == "nota":
            p = doc.add_paragraph()
            r = p.add_run(val)
            r.italic = True
            r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor.from_string(GRIS)
        elif tipo == "tabla":
            _tabla_docx(doc, val, ancho)
        elif tipo == "salto":
            doc.add_page_break()
    doc.core_properties.title = titulo_doc
    doc.core_properties.author = emp["nombre"] or "PRO-REFORMAS"
    salida = BytesIO()
    doc.save(salida)
    return salida.getvalue()


# ---------------------------------------------------------------- PDF (reportlab)
_SUSTITUCIONES_PDF = {"√": "raíz de ", "≥": ">=", "≤": "<=", "✔": "OK", "✘": "NO", "→": "->",
                      "≈": "~", "−": "-", "’": "'", "“": '"', "”": '"', "…": "..."}


def _limpiar_pdf(s) -> str:
    s = str(s)
    for a, b in _SUSTITUCIONES_PDF.items():
        s = s.replace(a, b)
    return s


def _pdf_txt(s) -> str:
    return xml_escape(_limpiar_pdf(s))


def render_pdf(bloques, d, titulo_doc) -> bytes:
    W, H = A4
    ml = mr = 2.0 * cm
    ancho = W - ml - mr
    azul, amar = colors.HexColor("#" + AZUL), colors.HexColor("#" + AMARILLO)
    gris, claro = colors.HexColor("#" + GRIS), colors.HexColor("#" + CLARO)
    emp = d["empresa"]
    nombre_emp = _limpiar_pdf(emp["nombre"] or "PRO-REFORMAS")
    linea_emp = _limpiar_pdf(" · ".join(x for x in (emp["nif"], emp["contacto"]) if x))
    titulo_doc = _limpiar_pdf(titulo_doc)
    ref_txt = _limpiar_pdf(f"Ref. {d['ref']} · {d['fecha']:%d/%m/%Y}")

    est = {
        "p": ParagraphStyle("p", fontName="Helvetica", fontSize=9.5, leading=13, spaceAfter=5),
        "h1": ParagraphStyle("h1", fontName="Helvetica-Bold", fontSize=13, leading=16, spaceBefore=12,
                             spaceAfter=5, textColor=azul),
        "h2": ParagraphStyle("h2", fontName="Helvetica-Bold", fontSize=10.5, leading=13, spaceBefore=8,
                             spaceAfter=4, textColor=azul),
        "nota": ParagraphStyle("nota", fontName="Helvetica-Oblique", fontSize=8, leading=10.5,
                               spaceAfter=6, textColor=gris),
        "titulo": ParagraphStyle("titulo", fontName="Helvetica-Bold", fontSize=18, leading=22,
                                 textColor=azul, spaceAfter=3),
        "sub": ParagraphStyle("sub", fontName="Helvetica-Oblique", fontSize=10.5, leading=14,
                              textColor=gris, spaceAfter=6),
        "lin": ParagraphStyle("lin", fontName="Helvetica", fontSize=9, leading=12),
        "td": ParagraphStyle("td", fontName="Helvetica", fontSize=8, leading=10),
        "tdr": ParagraphStyle("tdr", fontName="Helvetica", fontSize=8, leading=10, alignment=2),
        "th": ParagraphStyle("th", fontName="Helvetica-Bold", fontSize=8, leading=10,
                             textColor=colors.white),
        "thr": ParagraphStyle("thr", fontName="Helvetica-Bold", fontSize=8, leading=10,
                              textColor=colors.white, alignment=2),
    }

    class Lienzo(rl_canvas.Canvas):
        def __init__(self, *a, **k):
            super().__init__(*a, **k)
            self._estados = []

        def showPage(self):
            self._estados.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            total = len(self._estados)
            for estado in self._estados:
                self.__dict__.update(estado)
                self._decorar(total)
                super().showPage()
            super().save()

        def _decorar(self, total):
            self.saveState()
            self.setFillColor(azul)
            self.rect(0, H - 1.9 * cm, W, 1.9 * cm, stroke=0, fill=1)
            self.setFillColor(amar)
            self.rect(0, H - 2.02 * cm, W, 0.12 * cm, stroke=0, fill=1)
            self.setFillColor(colors.white)
            self.setFont("Helvetica-Bold", 12)
            self.drawString(ml, H - 0.95 * cm, nombre_emp)
            self.setFont("Helvetica", 8)
            self.drawString(ml, H - 1.45 * cm, linea_emp)
            self.setFont("Helvetica-Bold", 9)
            self.drawRightString(W - mr, H - 0.95 * cm, titulo_doc)
            self.setFont("Helvetica", 8)
            self.drawRightString(W - mr, H - 1.45 * cm, ref_txt)
            self.setStrokeColor(amar)
            self.setLineWidth(0.8)
            self.line(ml, 1.45 * cm, W - mr, 1.45 * cm)
            self.setFillColor(gris)
            self.setFont("Helvetica", 7)
            self.drawString(ml, 0.95 * cm, "Documento orientativo: no sustituye el proyecto ni la "
                                           "memoria técnica de un técnico competente.")
            self.setFont("Helvetica", 8)
            self.drawRightString(W - mr, 0.95 * cm, f"Página {self.getPageNumber()} de {total}")
            self.restoreState()

    def tabla_pdf(t):
        cols, rows = t["cols"], t["rows"]
        pesos = t["w"] or [1] * len(cols)
        tot = sum(pesos)
        anchos = [ancho * p / tot for p in pesos]
        num = set(t["num"])
        datos = [[Paragraph(_pdf_txt(c), est["thr"] if i in num else est["th"])
                  for i, c in enumerate(cols)]]
        n_res = t["resaltar"]
        for ri, fila in enumerate(rows):
            resaltada = n_res and ri >= len(rows) - n_res
            celdas = []
            for i, v in enumerate(fila):
                txt = _pdf_txt(v)
                if resaltada:
                    txt = f"<b>{txt}</b>"
                celdas.append(Paragraph(txt, est["tdr"] if i in num else est["td"]))
            datos.append(celdas)
        tb = Table(datos, colWidths=anchos, repeatRows=1)
        estilo = [("BACKGROUND", (0, 0), (-1, 0), azul),
                  ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#C9D6E2")),
                  ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, claro]),
                  ("VALIGN", (0, 0), (-1, -1), "TOP"),
                  ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                  ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4)]
        if n_res:
            estilo.append(("BACKGROUND", (0, -n_res), (-1, -1), colors.HexColor("#FFE9A6")))
        tb.setStyle(TableStyle(estilo))
        return tb

    elems = []
    for tipo, val in bloques:
        if tipo == "portada":
            elems.append(Paragraph(_pdf_txt(val["titulo"]), est["titulo"]))
            elems.append(Paragraph(_pdf_txt(val["subtitulo"]), est["sub"]))
            for linea in val["lineas"]:
                elems.append(Paragraph(_pdf_txt(linea), est["lin"]))
            elems.append(Spacer(1, 8))
        elif tipo in ("h1", "h2", "p", "nota"):
            if tipo in ("h1", "h2"):
                elems.append(CondPageBreak(4.5 * cm))
            elems.append(Paragraph(_pdf_txt(val), est[tipo]))
        elif tipo == "bullets":
            for item in val:
                elems.append(Paragraph(_pdf_txt(item), est["p"], bulletText="•"))
        elif tipo == "tabla":
            elems.append(tabla_pdf(val))
            elems.append(Spacer(1, 6))
        elif tipo == "salto":
            elems.append(PageBreak())

    buf = BytesIO()
    plantilla = SimpleDocTemplate(buf, pagesize=A4, leftMargin=ml, rightMargin=mr,
                                  topMargin=2.9 * cm, bottomMargin=2.0 * cm,
                                  title=titulo_doc, author=nombre_emp)
    plantilla.build(elems, canvasmaker=Lienzo)
    return buf.getvalue()


# ---------------------------------------------------------------- Excel (.xlsx)
EUR = '#,##0.00 "€"'
FUENTE = "Arial"


def generar_excel(d) -> bytes:
    wb = Workbook()
    emp, pres = d["empresa"], d["pres"]
    lado = Side(style="thin", color="C9D6E2")
    borde = Border(left=lado, right=lado, top=lado, bottom=lado)
    rell_azul = PatternFill("solid", fgColor=AZUL)
    rell_amar = PatternFill("solid", fgColor=AMARILLO)
    rell_claro = PatternFill("solid", fgColor=CLARO)

    def preparar(ws, titulo, ncols, apaisada=False):
        """Encabezado de hoja + cabecera/pie de impresión. Devuelve la primera fila libre."""
        letra = get_column_letter(ncols)
        for r in (1, 2):
            for c in range(1, ncols + 1):
                ws.cell(r, c).fill = rell_azul
        ws["A1"] = emp["nombre"] or "PRO-REFORMAS"
        ws["A1"].font = Font(name=FUENTE, size=14, bold=True, color="FFFFFF")
        ws["A2"] = " · ".join(x for x in (emp["nif"], emp["contacto"]) if x)
        ws["A2"].font = Font(name=FUENTE, size=9, color="CDEBFA")
        ws.cell(1, ncols, f"Ref. {d['ref']}").font = Font(name=FUENTE, size=10, bold=True,
                                                          color="FFFFFF")
        ws.cell(1, ncols).alignment = Alignment(horizontal="right")
        ws.cell(2, ncols, f"{d['fecha']:%d/%m/%Y}").font = Font(name=FUENTE, size=9, color="CDEBFA")
        ws.cell(2, ncols).alignment = Alignment(horizontal="right")
        ws["A3"] = titulo
        ws["A3"].font = Font(name=FUENTE, size=13, bold=True, color=AZUL)
        for i, (etq, val) in enumerate([("Cliente:", d["cliente"] or "Sin especificar"),
                                         ("Obra:", d["obra"] or "Sin especificar"),
                                         ("Superficie (m²):", d["superficie"]),
                                         ("Validez (días):", d["validez_dias"])], start=4):
            ws.cell(i, 1, etq).font = Font(name=FUENTE, bold=True)
            ws.cell(i, 2, val).font = Font(name=FUENTE)
            ws.cell(i, 2).alignment = Alignment(horizontal="left")
        ws.oddHeader.left.text = emp["nombre"] or "PRO-REFORMAS"
        ws.oddHeader.right.text = f"Ref. {d['ref']}"
        ws.oddFooter.left.text = "Documento orientativo"
        ws.oddFooter.right.text = "Página &P de &N"
        ws.page_setup.orientation = "landscape" if apaisada else "portrait"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.page_setup.fitToHeight = 0
        _ = letra
        return 10

    def cabecera_tabla(ws, fila, titulos):
        for c, t in enumerate(titulos, start=1):
            celda = ws.cell(fila, c, t)
            celda.font = Font(name=FUENTE, bold=True, color="FFFFFF")
            celda.fill = rell_azul
            celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            celda.border = borde

    def hoja_lineas(ws, titulo, filas, etq_concepto, con_obs):
        ncols = 7 if con_obs else 6
        fila0 = preparar(ws, titulo, ncols, apaisada=True)
        titulos = ["Capítulo", etq_concepto, "Ud", "Cantidad", "Precio unit. (€)", "Importe (€)"]
        if con_obs:
            titulos.append("Observaciones")
        cabecera_tabla(ws, fila0, titulos)
        primera = fila0 + 1
        for i, f in enumerate(filas):
            r = primera + i
            concepto = f.get("Material") or f.get("Partida")
            valores = [f["Capítulo"], concepto, f["Ud"], f["Cantidad"], f["Precio unit. (€)"]]
            for c, v in enumerate(valores, start=1):
                ws.cell(r, c, v)
            ws.cell(r, 4).number_format = "#,##0.00" if not con_obs else "#,##0"
            ws.cell(r, 5).number_format = EUR
            ws.cell(r, 6, f"=D{r}*E{r}").number_format = EUR
            if con_obs:
                ws.cell(r, 7, f.get("Nota", ""))
            for c in range(1, ncols + 1):
                celda = ws.cell(r, c)
                celda.font = Font(name=FUENTE, size=10)
                celda.border = borde
                celda.alignment = Alignment(vertical="top", wrap_text=(c in (2, 7)))
                if i % 2 == 1:
                    celda.fill = rell_claro
        ultima = primera + max(len(filas), 1) - 1
        t = ultima + 2
        ws.cell(t, 5, "TOTAL").font = Font(name=FUENTE, bold=True)
        ws.cell(t, 6, f"=SUM(F{primera}:F{ultima})").number_format = EUR
        ws.cell(t, 6).font = Font(name=FUENTE, bold=True)
        ws.cell(t, 6).fill = rell_amar
        for col, ancho in zip("ABCDEFG", (24, 62, 8, 12, 16, 16, 40)):
            ws.column_dimensions[col].width = ancho
        ws.freeze_panes = ws.cell(fila0 + 1, 1)
        ws.print_title_rows = f"{fila0}:{fila0}"
        return primera, ultima

    # --- Hoja 2 (se crea primero para conocer su rango): Mediciones y presupuesto
    ws_res = wb.active
    ws_res.title = "Resumen económico"
    ws_med = wb.create_sheet("Mediciones y presupuesto")
    materiales = pres["metodo"] == "materiales"
    filas_med = d["mats"] if materiales else pres["lineas"]
    prim, ult = hoja_lineas(ws_med, "Mediciones y presupuesto", filas_med,
                            "Material" if materiales else "Partida", materiales)
    rango_a = f"'Mediciones y presupuesto'!$A${prim}:$A${ult}"
    rango_f = f"'Mediciones y presupuesto'!$F${prim}:$F${ult}"

    # --- Hoja 1: Resumen económico
    ws = ws_res
    fila0 = preparar(ws, "Resumen económico", 4)
    ws.cell(fila0, 1, "IVA (%)").font = Font(name=FUENTE, bold=True)
    ws.cell(fila0, 2, pres["iva_pct"]).font = Font(name=FUENTE)
    ws.cell(fila0, 2).alignment = Alignment(horizontal="left")
    fila_iva = fila0
    fila_mo = None
    if materiales:
        ws.cell(fila0 + 1, 1, "Mano de obra (% sobre materiales)").font = Font(name=FUENTE, bold=True)
        ws.cell(fila0 + 1, 2, pres["mo_pct"]).font = Font(name=FUENTE)
        ws.cell(fila0 + 1, 2).alignment = Alignment(horizontal="left")
        fila_mo = fila0 + 1
    cab = fila0 + 3
    cabecera_tabla(ws, cab, ["Capítulo / módulo", "Materiales (€)", "Mano de obra (€)",
                              "Total capítulo (€)"])
    r = cab + 1
    primera_cap = r
    for i, cpt in enumerate(pres["caps"], start=1):
        ws.cell(r, 1, f"Capítulo {i}: {cpt['nombre']}")
        crit = f'"{cpt["cap"]}"'
        suma = f"SUMIF({rango_a},{crit},{rango_f})"
        if materiales:
            ws.cell(r, 2, f"={suma}")
            ws.cell(r, 3, f"=B{r}*$B${fila_mo}/100")
            ws.cell(r, 4, f"=B{r}+C{r}")
        else:
            ws.cell(r, 2, "—")
            ws.cell(r, 3, "—")
            ws.cell(r, 4, f"={suma}")
        for c in range(1, 5):
            celda = ws.cell(r, c)
            celda.font = Font(name=FUENTE)
            celda.border = borde
            if c > 1:
                celda.number_format = EUR
                celda.alignment = Alignment(horizontal="right")
        r += 1
    ultima_cap = r - 1
    r += 1
    filas_tot = [("Base imponible", f"=SUM(D{primera_cap}:D{ultima_cap})"),
                 ("IVA", f"=D{r}*$B${fila_iva}/100"),
                 ("TOTAL PRESUPUESTO", f"=D{r}+D{r + 1}")]
    if materiales:
        ws.cell(r, 2, f"=SUM(B{primera_cap}:B{ultima_cap})").number_format = EUR
        ws.cell(r, 3, f"=SUM(C{primera_cap}:C{ultima_cap})").number_format = EUR
    for k, (etq, formula) in enumerate(filas_tot):
        rr = r + k
        ws.cell(rr, 1, etq).font = Font(name=FUENTE, bold=True)
        ws.cell(rr, 4, formula).number_format = EUR
        ws.cell(rr, 4).font = Font(name=FUENTE, bold=True)
    for c in range(1, 5):
        ws.cell(r + 2, c).fill = rell_amar
    ws.cell(r + 4, 1, f"Presupuesto válido {d['validez_dias']} días naturales. "
                      f"Estimación orientativa con la normativa a {NORMATIVA_FECHA}; no sustituye "
                      "proyecto ni memoria técnica.").font = Font(name=FUENTE, italic=True, size=9)
    for col, ancho in zip("ABCD", (52, 18, 18, 20)):
        ws.column_dimensions[col].width = ancho

    # --- Hoja 3: listado de materiales (si el presupuesto es por tarifas)
    if not materiales and d["mats"]:
        hoja_lineas(wb.create_sheet("Listado de materiales"), "Listado de materiales", d["mats"],
                    "Material", True)

    # --- Hoja: circuitos y protecciones
    e = d["elec"]
    if d["inc"]["inc_elec"] and e and e["filas_circ"]:
        wc = wb.create_sheet("Circuitos y protecciones")
        f0 = preparar(wc, "Circuitos y protecciones (REBT)", 7, apaisada=True)
        cabecera_tabla(wc, f0, ["Circuito", "Descripción", "Nº circuitos", "PIA (A)",
                                 "Sección (mm²)", "Tubo (mm)", "Puntos previstos"])
        rr = f0 + 1
        for f in e["filas_circ"]:
            for c, v in enumerate([f["Circuito"], f["Descripción"], f["Nº circuitos"], f["PIA (A)"],
                                   f["Sección (mm²)"], f["Tubo (mm)"], f["Puntos previstos"]],
                                  start=1):
                wc.cell(rr, c, v).border = borde
            rr += 1
        rr += 1
        cabecera_tabla(wc, rr, ["Elemento", "Valor", "Observaciones"])
        for p in e["prot"]:
            rr += 1
            for c, v in enumerate([p["Elemento"], p["Valor"], p["Observaciones"]], start=1):
                wc.cell(rr, c, v).border = borde
        for col, ancho in zip("ABCDEFG", (30, 40, 14, 10, 14, 12, 16)):
            wc.column_dimensions[col].width = ancho

    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()


# ---------------------------------------------------------------- Biblioteca (utilidades)
def archivos_locales() -> dict:
    """PDF de la carpeta 'biblioteca' junto a app.py: {nombre_archivo: Path}."""
    if not CARPETA_BIBLIOTECA.is_dir():
        return {}
    return {p.name: p for p in sorted(CARPETA_BIBLIOTECA.glob("*.pdf"))}


def visor_pdf(contenido: bytes, alto: int = 650):
    """Muestra un PDF dentro de la app (st.pdf si existe; si no, un visor incrustado)."""
    try:
        st.pdf(contenido, height=alto)
        return
    except Exception:
        pass
    b64 = base64.b64encode(contenido).decode()
    st.markdown(f'<iframe src="data:application/pdf;base64,{b64}" width="100%" height="{alto}" '
                'style="border:1px solid #9BC4DA;border-radius:8px;"></iframe>',
                unsafe_allow_html=True)


st.set_page_config(page_title="PRO-REFORMAS | Electro·Hidro", page_icon="⚡", layout="wide")

ESTILO = """
<style>
.stApp{background:#F3F8FC;}
[data-testid="stMainBlockContainer"],.block-container{padding-top:4.5rem !important;max-width:1250px;}
header[data-testid="stHeader"]{background:#0B2545 !important;}
header[data-testid="stHeader"] :is(button,a,span,p){color:#EAF4FB !important;}
header[data-testid="stHeader"] svg{fill:#EAF4FB !important;}
[data-testid="stMain"] :is(p,li,h1,h2,h3,h4,label){color:#10243E;}
[data-testid="stAlert"] *{color:#10243E !important;}

/* Cabecera */
.cabecera{background:linear-gradient(90deg,#0B2545 0%,#0B5E7E 65%,#0096C7 100%);
  border-bottom:6px solid #FFC400;border-radius:12px;padding:1.2rem 1.8rem;margin-bottom:1rem;}
.cabecera .t{color:#FFFFFF !important;font-size:2rem;font-weight:800;line-height:1.2;}
.cabecera .s{color:#CDEBFA !important;font-size:1.02rem;margin-top:.3rem;}

/* Bandas de sección por capítulo */
.banda{border-radius:10px;padding:.5rem 1rem;font-weight:800;font-size:1.08rem;margin:.5rem 0 .6rem 0;
  display:flex;justify-content:space-between;align-items:center;}
.banda .etq{font-size:.76rem;font-weight:700;padding:.15rem .65rem;border-radius:20px;background:rgba(255,255,255,.28);}
.b-elec{background:#FFC400;color:#0B2545 !important;}
.b-hidro{background:#0096C7;color:#FFFFFF !important;}
.b-clima{background:#2EC4B6;color:#06312E !important;}
.b-teleco{background:#5B5BD6;color:#FFFFFF !important;}
.b-alba{background:#6B7A8F;color:#FFFFFF !important;}
.b-marino{background:#0B2545;color:#FFFFFF !important;}
.b-cctv{background:#D1495B;color:#FFFFFF !important;}

/* Tarjetas de resultado */
.kpi{background:#FFFFFF;border-radius:10px;padding:.65rem 1rem;border-left:7px solid var(--c);
  box-shadow:0 1px 4px rgba(11,37,69,.14);margin:.3rem 0;}
.kpi .l{font-size:.78rem;font-weight:700;color:#5B7088 !important;text-transform:uppercase;letter-spacing:.03em;}
.kpi .v{font-size:1.65rem;font-weight:800;color:#0B2545 !important;line-height:1.25;}
.kpi .n{font-size:.82rem;color:#5B7088 !important;}
.k-elec{--c:#FFC400;} .k-hidro{--c:#0096C7;} .k-clima{--c:#2EC4B6;} .k-marino{--c:#0B2545;} .k-cctv{--c:#D1495B;}
.kpi.total{background:#FFC400;}
.kpi.total .l,.kpi.total .v,.kpi.total .n{color:#0B2545 !important;}

/* Pestañas */
button[role="tab"] p{font-weight:700;font-size:1.02rem;}
button[role="tab"][aria-selected="true"]{border-bottom:4px solid #FFC400 !important;}
button[role="tab"][aria-selected="true"] p{color:#0B5E7E !important;}
div[data-baseweb="tab-highlight"]{background:#FFC400 !important;height:4px !important;}

/* Botones */
.stButton>button,.stDownloadButton>button{background:#FFC400;color:#0B2545 !important;border:none;
  font-weight:800;border-radius:10px;width:100% !important;}
.stButton>button:hover,.stDownloadButton>button:hover{background:#E6AF00;color:#000 !important;}

/* Campos */
.stApp :is(input,textarea){background:#FFFFFF !important;color:#10243E !important;
  -webkit-text-fill-color:#10243E !important;caret-color:#10243E;}
.stApp div:has(> input:not([type="checkbox"]):not([type="radio"])),
.stApp div:has(> div > input:not([type="checkbox"]):not([type="radio"])),
[data-baseweb="input"],[data-baseweb="base-input"],[data-baseweb="select"]>div{background:#FFFFFF !important;}
[data-testid="stNumberInputContainer"],[data-testid="stTextInputRootElement"],[data-baseweb="select"]>div{
  border:2px solid #9BC4DA !important;border-radius:8px !important;}
[data-baseweb="select"] *{color:#10243E !important;}
[data-testid="stNumberInputStepUp"],[data-testid="stNumberInputStepDown"]{background:#EAF4FB !important;color:#10243E !important;}
[data-testid="stNumberInputStepUp"] svg,[data-testid="stNumberInputStepDown"] svg{fill:#10243E !important;color:#10243E !important;}
[data-baseweb="popover"] :is(ul,li,[role="option"],[role="listbox"]){background:#FFFFFF !important;color:#10243E !important;}
[data-baseweb="popover"] li:hover{background:#EAF4FB !important;}

/* Tablas propias (siempre claras) */
.tabla-wrap{overflow-x:auto;margin:.4rem 0 .9rem 0;}
.tabla{width:100%;border-collapse:collapse;background:#FFFFFF;box-shadow:0 1px 4px rgba(11,37,69,.14);}
.tabla th{background:#0B2545;color:#FFFFFF !important;padding:.5rem .7rem;text-align:left;font-size:.85rem;}
.tabla td{background:#FFFFFF;color:#10243E !important;padding:.45rem .7rem;border-bottom:1px solid #E1ECF4;}
.tabla tr:nth-child(even) td{background:#F3F8FC;}

/* Barra lateral */
section[data-testid="stSidebar"]{background:#0B2545 !important;}
section[data-testid="stSidebar"] :is(p,span,label,h1,h2,h3,h4,li,div){color:#EAF4FB;}
section[data-testid="stSidebar"] :is([data-testid="stNumberInputContainer"],[data-testid="stTextInputRootElement"],[data-baseweb="select"]>div){border:2px solid #FFC400 !important;}
section[data-testid="stSidebar"] [data-testid="stSelectbox"] *{color:#10243E !important;}
.logo{font-size:1.45rem;font-weight:800;color:#FFC400 !important;margin-bottom:.4rem;}
.barra{background:rgba(255,255,255,.22);border-radius:20px;height:10px;margin:.2rem 0 .5rem 0;}
.barra>div{background:#FFC400;height:10px;border-radius:20px;}
</style>
<div class="cabecera">
  <div class="t">⚡ PRO-REFORMAS 💧</div>
  <div class="s">Electricidad y fontanería: cálculo técnico orientativo y presupuestos en 4 pasos</div>
</div>
"""
st.markdown(ESTILO, unsafe_allow_html=True)


def banda(clase, titulo, etiqueta=""):
    etq = f'<span class="etq">{etiqueta}</span>' if etiqueta else ""
    st.markdown(f'<div class="banda b-{clase}"><span>{titulo}</span>{etq}</div>', unsafe_allow_html=True)


def tabla(df, formatos=None):
    html = df.to_html(index=False, escape=True, border=0, classes="tabla", formatters=formatos or {})
    st.markdown(f'<div class="tabla-wrap">{html.replace(chr(10), "")}</div>', unsafe_allow_html=True)


def kpi(clase, etiqueta, valor, nota="", extra=""):
    st.markdown(f'<div class="kpi k-{clase} {extra}"><div class="l">{etiqueta}</div>'
                f'<div class="v">{valor}</div><div class="n">{nota}</div></div>', unsafe_allow_html=True)


# --- Estado: valores por defecto (las pestañas siempre se renderizan, así que no se pierden) ---
DEFAULTS = {
    "cliente_nombre": "", "cliente_obra": "", "superficie": 0.0, "tipo_iva": 21,
    "empresa_nombre": "", "empresa_nif": "", "empresa_contacto": "", "validez_dias": 30,
    "ref_manual": "",
    "inc_elec": True, "inc_font": True, "inc_clima": False, "inc_teleco": False, "inc_cctv": False,
    "n_vestibulos": 0, "n_dormitorios": 0, "n_salones": 0, "n_banos": 0, "n_aseos": 0,
    "n_cocinas": 0, "n_galerias": 0, "n_terrazas": 0, "n_trasteros": 0,
    "pasillo_m": 0.0, "sup_salon_m2": 0.0, "tomas_datos_dorm": 1, "tomas_datos_salon": 2,
    "grado": "Básica", "potencia_prevista": 5750, "longitud_di": 15.0,
    "aislamiento": "PVC (70 °C)", "suministro_di": SUMINISTRO_DI_DEFECTO,
    "calefaccion_elec": False, "aire_acond": False, "secadora": False, "incluir_spd": False,
    "velocidad_ms": 1.5, "aislamiento_termico": "Media (100 W/m²)",
    "cctv_tecnologia": CCTV_TECNOLOGIAS[0], "cctv_res": list(CCTV_RESOLUCIONES)[0],
    "cctv_dias": 30, "cctv_domo": 0, "cctv_bullet": 0, "cctv_ptz": 0, "cctv_dist_m": 20.0,
    "cctv_monitor": True,
    "m_punto": 6.0, "m_dedicado": 12.0, "m_toma_datos": 15.0, "desperdicio_pct": 10.0,
    "metodo_pres": "materiales", "mo_pct": 35.0, "docs": None, "docs_firma": "",
    "pu_elec_m2": 0.0, "pu_elec_cgmp": 0.0, "pu_elec_spd": 0.0, "pu_font_bano": 0.0,
    "pu_font_aseo": 0.0, "pu_font_cocina": 0.0, "pu_clima_kw": 0.0, "pu_teleco_toma": 0.0,
    "pu_cctv_camara": 0.0, "pu_cctv_ptz": 0.0, "pu_cctv_grabador": 0.0, "pu_cctv_monitor": 0.0,
}
for clave, valor in DEFAULTS.items():
    st.session_state.setdefault(clave, valor)
s = st.session_state

CLAVES_ESTANCIAS = ["n_vestibulos", "n_dormitorios", "n_salones", "n_banos", "n_aseos",
                    "n_cocinas", "n_galerias", "n_terrazas", "n_trasteros", "pasillo_m"]
PLANTILLAS = {
    "🏢 Piso 2 hab.": dict(superficie=70.0, n_vestibulos=1, n_dormitorios=2, n_salones=1, n_banos=1,
                          n_aseos=0, n_cocinas=1, n_galerias=0, n_terrazas=1, n_trasteros=0,
                          pasillo_m=5.0),
    "🏢 Piso 3 hab.": dict(superficie=90.0, n_vestibulos=1, n_dormitorios=3, n_salones=1, n_banos=1,
                          n_aseos=1, n_cocinas=1, n_galerias=1, n_terrazas=1, n_trasteros=0,
                          pasillo_m=8.0),
    "🏡 Casa 4 hab.": dict(superficie=160.0, n_vestibulos=1, n_dormitorios=4, n_salones=1, n_banos=2,
                          n_aseos=1, n_cocinas=1, n_galerias=1, n_terrazas=2, n_trasteros=1,
                          pasillo_m=12.0),
    "🧹 Limpiar": dict(superficie=0.0, n_vestibulos=0, n_dormitorios=0, n_salones=0, n_banos=0,
                      n_aseos=0, n_cocinas=0, n_galerias=0, n_terrazas=0, n_trasteros=0,
                      pasillo_m=0.0, sup_salon_m2=0.0),
}


def aplicar_plantilla(nombre):
    for clave, valor in PLANTILLAS[nombre].items():
        st.session_state[clave] = valor


n_estancias = sum(s[k] for k in CLAVES_ESTANCIAS if k.startswith("n_"))

# --- Barra lateral ---
with st.sidebar:
    st.markdown('<div class="logo">⚡ PRO-REFORMAS 💧</div>', unsafe_allow_html=True)
    st.markdown("### 👤 Cliente y obra")
    st.text_input("Nombre del cliente / empresa", key="cliente_nombre")
    st.text_input("Dirección de la obra", key="cliente_obra")
    st.number_input("Superficie útil afectada (m²)", min_value=0.0, step=5.0, key="superficie")
    with st.expander("🏢 Tu empresa (encabezado de los documentos)"):
        st.text_input("Nombre o razón social", key="empresa_nombre")
        st.text_input("NIF / CIF", key="empresa_nif")
        st.text_input("Teléfono y email", key="empresa_contacto")
        st.text_input("Nº de presupuesto (si lo dejas vacío, se genera por fecha)", key="ref_manual")
        st.number_input("Validez del presupuesto (días)", min_value=1, step=1, key="validez_dias")
    st.markdown("### 🔧 Qué incluye el presupuesto")
    st.checkbox("⚡ Electricidad", key="inc_elec")
    st.checkbox("💧 Fontanería", key="inc_font")
    st.checkbox("❄️ Climatización", key="inc_clima")
    st.checkbox("📡 Telecomunicaciones", key="inc_teleco")
    st.checkbox("📹 CCTV / Videovigilancia", key="inc_cctv")
    st.markdown("### 🏷️ IVA")
    st.selectbox("Tipo de IVA", [21, 10], key="tipo_iva", label_visibility="collapsed",
                 format_func=lambda x: f"{x}% ({'General' if x == 21 else 'Reducido reforma'})")
    st.caption("El 10 % solo aplica a reformas en vivienda de particulares, con más de 2 años de "
               "antigüedad y con materiales aportados por el contratista ≤ 40 % de la base. "
               "Confírmalo con tu gestor. Canarias, Ceuta y Melilla tienen otros impuestos.")
    progreso = st.container()   # se rellena al final, cuando ya hay presupuesto

t1, t2, t3, t4, t5 = st.tabs(["📋 1. Estancias", "🛠️ 2. Cálculos", "💶 3. Tarifas y materiales",
                              "📊 4. Presupuesto y descargas", "📚 Biblioteca"])

# --- PASO 1: ESTANCIAS ---
with t1:
    st.markdown("#### Empieza con una plantilla o rellena las estancias a mano")
    cols = st.columns(len(PLANTILLAS))
    for col, nombre in zip(cols, PLANTILLAS):
        col.button(nombre, on_click=aplicar_plantilla, args=(nombre,), key=f"plantilla_{nombre}")

    seca, humeda = st.columns(2)
    with seca:
        banda("elec", "⚡ ZONAS SECAS", "luz y enchufes")
        a, b = st.columns(2)
        a.number_input("🛌 Dormitorios", min_value=0, step=1, key="n_dormitorios")
        b.number_input("🛋️ Salones / comedores", min_value=0, step=1, key="n_salones")
        a.number_input("🚪 Vestíbulos / accesos", min_value=0, step=1, key="n_vestibulos")
        b.number_input("Superficie por salón (m²)", min_value=0.0, step=1.0, key="sup_salon_m2",
                       help="Opcional. Si la indicas, se ajustan los puntos de luz y tomas del salón.")
        a.number_input("🌿 Terrazas / balcones", min_value=0, step=1, key="n_terrazas")
        b.number_input("📦 Trasteros / garajes", min_value=0, step=1, key="n_trasteros")
        st.number_input("Longitud total de pasillos (m)", min_value=0.0, step=1.0, key="pasillo_m")
    with humeda:
        banda("hidro", "💧 ZONAS HÚMEDAS", "agua y desagües")
        a, b = st.columns(2)
        a.number_input("🛁 Baños (con ducha)", min_value=0, step=1, key="n_banos")
        b.number_input("🧻 Aseos (sin ducha)", min_value=0, step=1, key="n_aseos")
        a.number_input("🍳 Cocinas", min_value=0, step=1, key="n_cocinas")
        b.number_input("🧺 Galerías / lavaderos", min_value=0, step=1, key="n_galerias")

    banda("teleco", "📡 RED DE DATOS", "opcional")
    a, b, _ = st.columns(3)
    a.number_input("Tomas de datos por dormitorio", min_value=0, step=1, key="tomas_datos_dorm")
    b.number_input("Tomas de datos por salón", min_value=0, step=1, key="tomas_datos_salon")

    st.info(f"**{n_estancias} estancias** · **{s.superficie:.0f} m²** afectados. "
            "Recuerda indicar la superficie en la barra lateral.")
    st.caption("👉 Siguiente: pestaña **2. Cálculos**")

# --- PASO 2: CÁLCULOS ---
di = None
filas_circ, prot_rows = [], []
mec = {"puntos_luz": 0, "tomas_c2": 0, "tomas_c5": 0, "c7_por_superficie": False}
h = cl = cc = None
limite_caida = CAIDA_DI_PCT[s.suministro_di]
with t2:
    st.warning("Resultados orientativos. No sustituyen la memoria técnica ni el proyecto de un "
               "instalador autorizado o técnico competente.")
    st.caption(f"Normativa de referencia (estado a {NORMATIVA_FECHA}): {NORMATIVA_APLICADA}. "
               "La revisión del REBT no estaba publicada en el BOE a esa fecha y hay "
               "modificaciones del RITE y del CTE anunciadas: comprueba el BOE antes de cada proyecto.")

    # ELECTRICIDAD
    banda("elec", "⚡ ELECTRICIDAD · REBT", "En el presupuesto" if s.inc_elec else "No incluida")
    e1, e2 = st.columns(2)
    with e1:
        st.checkbox("Calefacción eléctrica (C8)", key="calefaccion_elec")
        st.checkbox("Aire acondicionado (C9)", key="aire_acond")
        st.checkbox("Secadora (C10)", key="secadora")
        st.checkbox("Protector contra sobretensiones en el cuadro", key="incluir_spd")
    recomendado = grado_recomendado(s.superficie, s.calefaccion_elec, s.aire_acond)
    with e2:
        st.selectbox("Grado de electrificación", list(POTENCIAS_PREVISTAS_W), key="grado")
        opciones = POTENCIAS_PREVISTAS_W[s.grado]
        if s.potencia_prevista not in opciones:
            s.potencia_prevista = opciones[0]
        st.selectbox("Potencia prevista (W)", opciones, key="potencia_prevista")
    if recomendado == "Elevada" and s.grado == "Básica":
        st.warning("Con más de 160 m², calefacción eléctrica o aire acondicionado corresponde "
                   "electrificación elevada.")
    if not s.incluir_spd:
        st.info("El borrador del nuevo REBT (sin publicar a " + NORMATIVA_FECHA + ") haría "
                "obligatoria la protección contra sobretensiones; hoy depende de la instalación "
                "(ITC-BT-23).")
    d1, d2, d3 = st.columns(3)
    d1.number_input("Longitud de la derivación individual (m)", min_value=1.0, step=1.0, key="longitud_di")
    d2.selectbox("Aislamiento del cable", list(CONDUCTIVIDAD_CU), key="aislamiento")
    d3.selectbox("Suministro (límite de caída de tensión)", list(CAIDA_DI_PCT), key="suministro_di")
    st.caption("En una reforma la derivación individual suele estar ya instalada: usa este cálculo "
               "para comprobar si la existente es suficiente.")
    limite_caida = CAIDA_DI_PCT[s.suministro_di]
    try:
        di = derivacion_individual(s.potencia_prevista, s.longitud_di, s.aislamiento,
                                   caida_max_pct=limite_caida)
        k1, k2, k3 = st.columns(3)
        with k1:
            kpi("elec", "Intensidad", f"{di['intensidad_a']:.1f} A")
        with k2:
            kpi("elec", "Sección derivación (Cu)", f"{di['seccion_mm2']:g} mm²",
                f"Caída {di['s_caida']:.2f} · intensidad {di['s_int']:g} · mínimo 6")
        with k3:
            limite_txt = f"{limite_caida:g}".replace(".", ",")
            kpi("elec", "Caída de tensión", f"{di['caida_pct']:.2f} %",
                f"Límite {limite_txt} % {'✔' if di['cumple'] else '✘'}")
    except ValueError as err:
        st.error(str(err))

    estancias = {"vestibulos": s.n_vestibulos, "dormitorios": s.n_dormitorios,
                 "salones": s.n_salones, "banos": s.n_banos, "aseos": s.n_aseos,
                 "cocinas": s.n_cocinas, "galerias": s.n_galerias,
                 "terrazas": s.n_terrazas, "trasteros": s.n_trasteros}
    filas_circ, mec = circuitos(
        estancias, s.pasillo_m, {"C8": s.calefaccion_elec, "C9": s.aire_acond, "C10": s.secadora},
        sup_salon_m2=s.sup_salon_m2, superficie_m2=s.superficie, grado=s.grado)
    st.markdown("**Circuitos y puntos mínimos** (ITC-BT-25, simplificado)")
    if filas_circ:
        tabla(pd.DataFrame(filas_circ))
        st.caption(f"Puntos de luz: {mec['puntos_luz']} · tomas C2: {mec['tomas_c2']} · "
                   f"tomas C5: {mec['tomas_c5']}. Si indicas la superficie del salón, sus mínimos "
                   "se ajustan; en pasillos se calculan por longitud. Revisa el resto con la tabla 2 "
                   "de la ITC-BT-25.")
        if mec["c7_por_superficie"]:
            st.caption("Con electrificación elevada y más de 160 m² se prevé un circuito C2 "
                       "adicional (C7).")
        st.markdown("**Protecciones del cuadro** (orientativo)")
        n_circuitos_total = sum(f["Nº circuitos"] for f in filas_circ)
        prot_rows = protecciones(s.potencia_prevista, n_circuitos_total, s.incluir_spd)
        tabla(pd.DataFrame(prot_rows))
        st.caption("Confirma calibres y número de diferenciales con la ITC-BT-25 y el esquema "
                   "real del cuadro.")
    else:
        st.info("Añade estancias en el Paso 1 para ver los circuitos.")

    # FONTANERÍA
    banda("hidro", "💧 FONTANERÍA · CTE DB-HS4", "En el presupuesto" if s.inc_font else "No incluida")
    st.number_input("Velocidad de cálculo del agua (m/s)", min_value=0.5, max_value=2.0, step=0.1,
                    key="velocidad_ms")
    h = hidraulica(s.n_banos, s.n_aseos, s.n_cocinas, s.n_galerias, s.velocidad_ms)
    if h["n_aparatos"]:
        k1, k2, k3, k4 = st.columns(4)
        with k1:
            kpi("hidro", "Aparatos", h["n_aparatos"])
        with k2:
            kpi("hidro", "Caudal instalado", f"{h['q_fria']:.2f} L/s", "Agua fría")
        with k3:
            kpi("hidro", "Caudal de diseño", f"{h['q_diseno']:.2f} L/s", f"Simultaneidad {h['kn']:.2f}")
        with k4:
            kpi("hidro", "Ø interior mínimo", f"{h['diametro_mm']:.1f} mm", "Tubería general")
        st.caption("Caudales por aparato según tabla 2.1 del DB-HS4. Falta dimensionar los ramales y "
                   "comprobar presión (mín. 100 kPa en grifos, 150 kPa en fluxores y calentadores; "
                   "máx. 500 kPa).")
    else:
        st.info("Añade baños, aseos o cocinas en el Paso 1.")

    # CCTV
    banda("cctv", "📹 CCTV · VIDEOVIGILANCIA", "En el presupuesto" if s.inc_cctv else "No incluida")
    v1, v2, v3 = st.columns(3)
    v1.selectbox("Tecnología", CCTV_TECNOLOGIAS, key="cctv_tecnologia")
    v2.selectbox("Resolución y caudal por cámara", list(CCTV_RESOLUCIONES), key="cctv_res")
    v3.number_input("Días de grabación", min_value=1, step=1, key="cctv_dias")
    w1, w2, w3, w4 = st.columns(4)
    w1.number_input("Cámaras domo (interior / techo)", min_value=0, step=1, key="cctv_domo")
    w2.number_input("Cámaras bullet (pared / exterior)", min_value=0, step=1, key="cctv_bullet")
    w3.number_input("Cámaras PTZ (360°)", min_value=0, step=1, key="cctv_ptz")
    w4.number_input("Distancia media por cámara (m)", min_value=1.0, step=1.0, key="cctv_dist_m")
    st.checkbox("Incluir monitor de visualización", key="cctv_monitor")
    cc = cctv(s.cctv_domo, s.cctv_bullet, s.cctv_ptz, s.cctv_tecnologia,
              CCTV_RESOLUCIONES[s.cctv_res], s.cctv_dias, s.cctv_dist_m, s.cctv_monitor)
    if cc["n_cam"] > 0:
        k1, k2, k3, k4 = st.columns(4)
        with k1:
            kpi("cctv", "Cámaras", cc["n_cam"],
                f"{cc['n_domo']} domo · {cc['n_bullet']} bullet · {cc['n_ptz']} PTZ")
        with k2:
            kpi("cctv", "Grabador", f"{cc['canales']} canales",
                f"{'NVR' if cc['ip'] else 'DVR'} × {cc['n_grabadores']}")
        with k3:
            kpi("cctv", "Almacenamiento", f"{cc['tb']:.1f} TB",
                f"{cc['n_discos']} × disco de {cc['disco_tb']} TB · {cc['gb_dia']:.0f} GB/día")
        with k4:
            if cc["ip"]:
                kpi("cctv", "Potencia PoE", f"{cc['poe_total_w']:.0f} W",
                    f"Switch PoE+ de {cc['puertos_poe']} puertos")
            else:
                kpi("cctv", "Alimentación 12 V", f"{cc['fuentes']} fuente(s)", "Multisalida de 8")
        st.caption(f"Cableado: {cc['cable_m']} m. Valores orientativos (H.265, grabación continua, "
                   "10 % de margen): comprueba las bahías de disco y el consumo real del grabador. "
                   "Pon el cartel de zona videovigilada y cumple el RGPD; si el sistema va a una "
                   "central receptora de alarmas, la instalación debe hacerla una empresa de "
                   "seguridad habilitada (Ley 5/2014).")
    else:
        st.info("Indica cuántas cámaras quieres para dimensionar el grabador y el almacenamiento.")

    # CLIMATIZACIÓN
    banda("clima", "❄️ CLIMATIZACIÓN · estimación", "En el presupuesto" if s.inc_clima else "No incluida")
    if s.superficie > 0:
        st.selectbox("Aislamiento térmico del inmueble", list(RATIOS_W_M2), key="aislamiento_termico")
        cl = climatizacion(s.superficie, calc_w_m2(s))
        k1, k2, _ = st.columns(3)
        with k1:
            kpi("clima", "Potencia térmica", f"{cl['kw']:.2f} kW")
        with k2:
            kpi("clima", "Equivalente", f"{cl['kcal_h']:.0f} kcal/h")
        st.caption("Estimación por ratio: no es un cálculo de cargas según RITE.")
    else:
        st.info("Introduce la superficie en la barra lateral.")
    st.caption("👉 Siguiente: pestaña **3. Tarifas y materiales**")

# --- Cantidades de materiales (según lo calculado) ---
tomas_datos = s.n_dormitorios * s.tomas_datos_dorm + s.n_salones * s.tomas_datos_salon
ctx = dict(
    inc_elec=s.inc_elec, inc_font=s.inc_font, inc_teleco=s.inc_teleco, inc_cctv=s.inc_cctv,
    inc_clima=s.inc_clima, superficie=s.superficie, n_estancias=n_estancias,
    n_dormitorios=s.n_dormitorios, n_salones=s.n_salones, n_banos=s.n_banos, n_aseos=s.n_aseos,
    n_cocinas=s.n_cocinas, n_galerias=s.n_galerias, filas_circ=filas_circ, mec=mec,
    potencia=s.potencia_prevista, spd=s.incluir_spd, di=di, longitud_di=s.longitud_di,
    aparatos=aparatos_contados(s.n_banos, s.n_aseos, s.n_cocinas, s.n_galerias),
    tomas_datos=tomas_datos, cctv=cc, cl=cl, m_punto=s.m_punto, m_dedicado=s.m_dedicado,
    m_toma_datos=s.m_toma_datos, desperdicio_pct=s.desperdicio_pct)
cantidades = calcular_materiales(ctx)

# --- PASO 3: TARIFAS Y MATERIALES ---
EMOJI_CAP = {CAP_ELEC: "⚡", CAP_FONT: "💧", CAP_TELE: "📡", CAP_CCTV: "📹", CAP_CLIMA: "❄️"}
precios_mat: dict = {}
with t3:
    st.markdown("#### Pon tus precios (sin IVA). Todos parten de 0 €")
    sub1, sub2, sub3 = st.tabs(["💶 Tarifas por partida", "🧰 Precios de materiales",
                                "📦 Listado de materiales de la obra"])
    with sub1:
        st.caption("Para presupuestar por superficie y unidades (€/m², €/ud).")
        a, b = st.columns(2)
        with a:
            banda("elec", "⚡ ELECTRICIDAD")
            st.number_input("Instalación eléctrica general (€/m²)", min_value=0.0, step=5.0, key="pu_elec_m2")
            st.number_input("Cuadro general y protecciones (€/ud)", min_value=0.0, step=25.0, key="pu_elec_cgmp")
            st.number_input("Protector contra sobretensiones (€/ud)", min_value=0.0, step=10.0, key="pu_elec_spd")
            banda("hidro", "💧 FONTANERÍA")
            st.number_input("Baño completo (€/ud)", min_value=0.0, step=50.0, key="pu_font_bano")
            st.number_input("Aseo (€/ud)", min_value=0.0, step=50.0, key="pu_font_aseo")
            st.number_input("Cocina (€/ud)", min_value=0.0, step=50.0, key="pu_font_cocina")
        with b:
            banda("cctv", "📹 CCTV / VIDEOVIGILANCIA")
            st.number_input("Cámara fija instalada (€/ud)", min_value=0.0, step=10.0, key="pu_cctv_camara")
            st.number_input("Cámara PTZ instalada (€/ud)", min_value=0.0, step=10.0, key="pu_cctv_ptz")
            st.number_input("Grabador con disco duro (€/ud)", min_value=0.0, step=10.0, key="pu_cctv_grabador")
            st.number_input("Monitor de visualización (€/ud)", min_value=0.0, step=10.0, key="pu_cctv_monitor")
            banda("clima", "❄️ CLIMATIZACIÓN")
            st.number_input("Instalación (€/kW)", min_value=0.0, step=25.0, key="pu_clima_kw")
            banda("teleco", "📡 TELECOMUNICACIONES")
            st.number_input("Toma de datos RJ45 Cat6 (€/toma)", min_value=0.0, step=5.0, key="pu_teleco_toma")

    with sub2:
        st.caption("Para presupuestar con materiales + mano de obra. Pon el precio de lo que uses: "
                   "el listado de la obra aplica estos precios a las cantidades calculadas.")
        with st.expander("⚙️ Parámetros de obra para estimar cantidades"):
            p1, p2, p3, p4 = st.columns(4)
            p1.number_input("Recorrido de línea por punto de luz/toma (m)", min_value=1.0, step=0.5,
                            key="m_punto")
            p2.number_input("Recorrido de circuitos dedicados (m)", min_value=1.0, step=1.0,
                            key="m_dedicado")
            p3.number_input("Recorrido por toma de datos (m)", min_value=1.0, step=1.0,
                            key="m_toma_datos")
            p4.number_input("Desperdicio de cable y tubo (%)", min_value=0.0, step=1.0,
                            key="desperdicio_pct")
        for cap, clave, nombre in CAPITULOS:
            items = [(k, n, u) for k, c_, n, u in CATALOGO if c_ == cap]
            base_df = pd.DataFrame(
                {"Material": [n for _, n, _ in items], "Ud": [u for _, _, u in items],
                 "Precio unit. (€)": [0.0] * len(items)}, index=[k for k, _, _ in items])
            with st.expander(f"{EMOJI_CAP[cap]} {nombre} · {len(items)} materiales"):
                editado = st.data_editor(
                    base_df, key=f"ed_{cap}", hide_index=True,
                    disabled=["Material", "Ud"],
                    column_config={
                        "Material": st.column_config.TextColumn("Material", width="large"),
                        "Ud": st.column_config.TextColumn("Ud", width="small"),
                        "Precio unit. (€)": st.column_config.NumberColumn(
                            "Precio unit. (€)", min_value=0.0, step=0.01, format="%.2f")})
                precios_mat.update(editado["Precio unit. (€)"].fillna(0.0).to_dict())

# --- Listado, presupuesto y datos de los documentos ---
mats = listado_materiales(cantidades, precios_mat)
hay_precios = (any(s[k] > 0 for k in DEFAULTS if k.startswith("pu_"))
               or any(f["Precio unit. (€)"] > 0 for f in mats))
lineas_tarifas = partidas(s)
incluidos = {clave: s[clave] for _, clave, _ in CAPITULOS}
pres = armar_presupuesto(lineas_tarifas, mats, s.metodo_pres, s.mo_pct, s.tipo_iva, incluidos)
importe_base = pres["base"]

with t3:
    with sub3:
        if not mats:
            st.info("Todavía no hay materiales: añade estancias o configura CCTV, y marca las "
                    "especialidades en la barra lateral.")
        for cap, clave, nombre in CAPITULOS:
            filas_cap = [f for f in mats if f["Capítulo"] == cap]
            if not filas_cap:
                continue
            banda({"Electricidad": "elec", "Fontanería": "hidro", "Telecomunicaciones": "teleco",
                   "CCTV y videovigilancia": "cctv", "Climatización": "clima"}[cap], nombre.upper(),
                  f"{len(filas_cap)} materiales")
            df = pd.DataFrame(filas_cap)[["Material", "Ud", "Cantidad", "Precio unit. (€)",
                                          "Importe (€)", "Nota"]]
            tabla(df, {"Cantidad": lambda v: f"{v:,.0f}", "Precio unit. (€)": lambda v: f"{v:,.2f} €",
                       "Importe (€)": lambda v: f"{v:,.2f} €"})
        if mats:
            st.caption("Cantidades orientativas: dependen de los parámetros de obra (recorridos y "
                       "desperdicio). Revísalas antes de comprar.")

datos = {
    "empresa": {"nombre": s.empresa_nombre.strip(), "nif": s.empresa_nif.strip(),
                "contacto": s.empresa_contacto.strip()},
    "cliente": s.cliente_nombre.strip(), "obra": s.cliente_obra.strip(), "fecha": date.today(),
    "ref": s.ref_manual.strip() or f"PR-{date.today():%Y%m%d}",
    "validez_dias": int(s.validez_dias), "superficie": s.superficie,
    "inc": {clave: bool(s[clave]) for _, clave, _ in CAPITULOS},
    "estancias": [("Vestíbulos / accesos", s.n_vestibulos), ("Dormitorios", s.n_dormitorios),
                  ("Salones / comedores", s.n_salones), ("Baños", s.n_banos), ("Aseos", s.n_aseos),
                  ("Cocinas", s.n_cocinas), ("Galerías / lavaderos", s.n_galerias),
                  ("Terrazas / balcones", s.n_terrazas), ("Trasteros / garajes", s.n_trasteros)],
    "elec": {"grado": s.grado, "potencia": s.potencia_prevista,
             "iga": next((a for a in IGA_ESTANDAR_A if a >= s.potencia_prevista / TENSION_V - 1e-6), None),
             "di": di, "longitud_di": s.longitud_di, "aislamiento": s.aislamiento,
             "suministro": s.suministro_di, "limite": limite_caida, "filas_circ": filas_circ,
             "prot": prot_rows, "spd": s.incluir_spd} if s.inc_elec else None,
    "font": {"hid": h, "velocidad": s.velocidad_ms} if s.inc_font and h else None,
    "tele": {"tomas": tomas_datos, "td_dorm": s.tomas_datos_dorm,
             "td_salon": s.tomas_datos_salon} if s.inc_teleco else None,
    "cctv": cc if s.inc_cctv else None,
    "clima": {"cl": cl, "ratio": s.aislamiento_termico} if s.inc_clima and cl else None,
    "pres": pres, "mats": mats,
    "precios_en_uso": any(f["Precio unit. (€)"] > 0 for f in mats),
}

with progreso:
    pasos = [("Datos de la obra", bool(s.cliente_nombre.strip()) and s.superficie > 0),
             ("Estancias", n_estancias > 0),
             ("Precios", hay_precios),
             ("Presupuesto con importe", importe_base > 0)]
    hechos = sum(ok for _, ok in pasos)
    st.markdown("### ✅ Progreso")
    st.markdown(f'<div class="barra"><div style="width:{hechos * 25}%"></div></div>', unsafe_allow_html=True)
    for nombre, ok in pasos:
        st.markdown(f"{'✅' if ok else '⬜'} {nombre}")

# --- PASO 4: PRESUPUESTO Y DESCARGAS ---
with t4:
    if not pres["caps"]:
        st.info("Marca al menos una especialidad en la barra lateral.")
    else:
        m1, m2 = st.columns([3, 1])
        m1.radio("Cómo quieres presupuestar", ["materiales", "tarifas"], key="metodo_pres",
                 horizontal=True,
                 format_func=lambda x: ("Materiales + mano de obra" if x == "materiales"
                                        else "Por partidas (tarifas €/m² y €/ud)"))
        m2.number_input("Mano de obra (% sobre materiales)", min_value=0.0, step=1.0, key="mo_pct",
                        disabled=(s.metodo_pres != "materiales"))
        if not hay_precios:
            st.warning("Aún no has puesto precios. Ve a la pestaña **3. Tarifas y materiales**.")
        if s.superficie <= 0 and n_estancias == 0:
            st.warning("Faltan la superficie y las estancias. Empieza en la pestaña **1. Estancias**.")

        if pres["metodo"] == "materiales":
            filas_vista = [{"Capítulo": c["nombre"], "Materiales (€)": c["materiales"],
                            "Mano de obra (€)": c["mano_obra"], "Total capítulo (€)": c["total"]}
                           for c in pres["caps"]]
            tabla(pd.DataFrame(filas_vista), {k: (lambda v: f"{v:,.2f} €") for k in
                                              ("Materiales (€)", "Mano de obra (€)",
                                               "Total capítulo (€)")})
        else:
            tabla(pd.DataFrame(pres["lineas"]), {"Cantidad": lambda v: f"{v:,.2f}",
                                                 "Precio unit. (€)": lambda v: f"{v:,.2f} €",
                                                 "Importe (€)": lambda v: f"{v:,.2f} €"})
        r1, r2, r3 = st.columns(3)
        with r1:
            kpi("marino", "Base imponible", f"{pres['base']:,.2f} €")
        with r2:
            kpi("hidro", f"IVA ({s.tipo_iva} %)", f"{pres['iva']:,.2f} €")
        with r3:
            kpi("elec", "TOTAL", f"{pres['total']:,.2f} €", extra="total")

        banda("marino", "📥 DESCARGAS", "Word · PDF · Excel")
        faltan = [n for n, ok in (("openpyxl (Excel)", HAY_OPENPYXL), ("python-docx (Word)", HAY_DOCX),
                                  ("reportlab (PDF)", HAY_REPORTLAB)) if not ok]
        if faltan:
            st.warning("Añade estas líneas a requirements.txt: " + ", ".join(faltan) + ".")
        firma = hashlib.md5(repr(datos).encode("utf-8")).hexdigest()
        if st.button("📄 Preparar documentos para descargar", key="preparar_docs"):
            with st.spinner("Generando Word, PDF y Excel..."):
                docs = {}
                if HAY_OPENPYXL:
                    docs["xlsx"] = generar_excel(datos)
                if HAY_DOCX:
                    docs["docx"] = render_docx(bloques_memoria(datos), datos,
                                               "Memoria técnica y presupuesto")
                if HAY_REPORTLAB:
                    docs["pdf_memoria"] = render_pdf(bloques_memoria(datos), datos,
                                                     "Memoria técnica y presupuesto")
                    docs["pdf_presupuesto"] = render_pdf(bloques_presupuesto(datos), datos,
                                                         "Presupuesto de reforma")
                    docs["pdf_materiales"] = render_pdf(
                        [("portada", {"titulo": "LISTADO DE MATERIALES",
                                      "subtitulo": "Materiales necesarios para la obra",
                                      "lineas": [f"Cliente: {datos['cliente'] or 'Sin especificar'}",
                                                 f"Obra: {datos['obra'] or 'Sin especificar'}",
                                                 f"Fecha: {datos['fecha']:%d/%m/%Y}   ·   "
                                                 f"Ref.: {datos['ref']}"]})]
                        + bloques_listado(datos), datos, "Listado de materiales")
                s.docs, s.docs_firma = docs, firma
        if s.docs and s.docs_firma != firma:
            st.warning("Has cambiado datos desde la última vez: pulsa **Preparar documentos** "
                       "otra vez para que los archivos reflejen los cambios.")
        elif s.docs:
            nombre = re.sub(r"[^\w\-]+", "_", s.cliente_nombre.strip()) or "Reforma"
            d1, d2, d3 = st.columns(3)
            if "docx" in s.docs:
                d1.download_button("📝 Memoria técnica y presupuesto (Word)", s.docs["docx"],
                                   file_name=f"Memoria_tecnica_{nombre}.docx", key="dl_docx",
                                   mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
            if "pdf_memoria" in s.docs:
                d2.download_button("📕 Memoria técnica y presupuesto (PDF)", s.docs["pdf_memoria"],
                                   file_name=f"Memoria_tecnica_{nombre}.pdf", key="dl_pdf_mem",
                                   mime="application/pdf")
                d3.download_button("📕 Presupuesto para el cliente (PDF)", s.docs["pdf_presupuesto"],
                                   file_name=f"Presupuesto_{nombre}.pdf", key="dl_pdf_pres",
                                   mime="application/pdf")
            e1_, e2_, _ = st.columns(3)
            if "xlsx" in s.docs:
                e1_.download_button("🟢 Presupuesto completo (Excel)", s.docs["xlsx"],
                                    file_name=f"Presupuesto_{nombre}.xlsx", key="dl_xlsx",
                                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            if "pdf_materiales" in s.docs:
                e2_.download_button("📦 Listado de materiales (PDF)", s.docs["pdf_materiales"],
                                    file_name=f"Listado_materiales_{nombre}.pdf", key="dl_pdf_mat",
                                    mime="application/pdf")
        st.caption("Los documentos llevan el encabezado con los datos de tu empresa (barra lateral).")

# --- BIBLIOTECA ---
with t5:
    n0, n1, n2, n3, n4, n5 = st.tabs(["📄 Normas y manuales", "⚡ REBT", "💧 CTE DB-HS", "❄️ RITE",
                                      "📡 Telecomunicaciones", "📹 CCTV"])
    with n0:
        locales = archivos_locales()
        catalogados = {b["archivo"] for b in BIBLIOTECA if b["archivo"]}
        extras = {n: p for n, p in locales.items() if n not in catalogados}
        q1, q2 = st.columns([2, 1])
        texto = q1.text_input("Buscar", key="bib_buscar", placeholder="REBT, videovigilancia, agua...")
        ambitos = ["Todos"] + sorted({b["ambito"] for b in BIBLIOTECA}) + (["Mis documentos"] if extras else [])
        ambito = q2.selectbox("Ámbito", ambitos, key="bib_ambito")
        st.caption("«Ver» abre la fuente oficial. Si copias los PDF a la carpeta **biblioteca/** del "
                   "repositorio, aquí podrás verlos y descargarlos sin salir de la app (también "
                   "aparecerán tus propios manuales).")

        def tarjeta(clave, titulo, sub, ver=None, pdf_url=None, ruta=None):
            with st.container(border=True):
                st.markdown(f"**{titulo}**  \n{sub}")
                ca, cb, cc_ = st.columns(3)
                if ver:
                    ca.link_button("👁️ Ver (fuente oficial)", ver)
                if ruta:
                    contenido = ruta.read_bytes()
                    cb.download_button("⬇️ Descargar PDF", contenido, file_name=ruta.name,
                                       mime="application/pdf", key=f"dl_{clave}")
                    if cc_.toggle("Ver aquí", key=f"ver_{clave}"):
                        visor_pdf(contenido)
                elif pdf_url:
                    cb.link_button("⬇️ Descargar PDF (oficial)", pdf_url)
                else:
                    cb.caption("Solo consulta en línea")

        mostrados = 0
        for b in BIBLIOTECA:
            if ambito not in ("Todos", b["ambito"]):
                continue
            if texto.strip() and texto.strip().lower() not in (b["titulo"] + b["desc"] + b["ambito"]).lower():
                continue
            tarjeta(b["id"], b["titulo"], f"{b['tipo']} · {b['ambito']} — {b['desc']}",
                    ver=b["ver"], pdf_url=b["pdf"], ruta=locales.get(b["archivo"]) if b["archivo"] else None)
            mostrados += 1
        if ambito in ("Todos", "Mis documentos"):
            for nombre_arch, ruta in extras.items():
                if texto.strip() and texto.strip().lower() not in nombre_arch.lower():
                    continue
                tarjeta("extra_" + nombre_arch, nombre_arch, "Documento propio · carpeta biblioteca/",
                        ruta=ruta)
                mostrados += 1
        if not mostrados:
            st.info("No hay documentos que coincidan con la búsqueda.")
        with st.expander("Nombres de archivo esperados en la carpeta biblioteca/"):
            st.markdown("\n".join(f"- `{b['archivo']}` — {b['titulo']}" for b in BIBLIOTECA if b["archivo"]))
    with n1:
        st.markdown("- **ITC-BT-25:** circuitos C1 (iluminación), C2 (tomas generales), C3 (cocina y "
                    "horno), C4 (lavadora, lavavajillas, termo), C5 (baños y auxiliares de cocina); "
                    "en electrificación elevada, C8 a C12, y un C2 adicional (C7) con más de 160 m².\n"
                    "- **ITC-BT-15:** derivación individual de 6 mm² como mínimo.\n"
                    "- **Caída de tensión en la derivación individual (ITC-BT-15):** 0,5 % con "
                    "contadores en más de un lugar, 1 % con contadores totalmente concentrados y "
                    "1,5 % para un único usuario sin línea general de alimentación. En la "
                    "instalación interior de viviendas, 3 % (ITC-BT-19).\n"
                    f"- **Estado a {NORMATIVA_FECHA}:** vigente el RD 842/2002. La revisión del "
                    "REBT (diferenciales tipo AC, sobretensiones, ITC-BT-53) no estaba publicada "
                    "en el BOE.")
    with n2:
        st.markdown("- **DB-HS4, tabla 2.1:** lavabo 0,10 · ducha 0,20 · fregadero 0,20 · "
                    "lavavajillas 0,15 · lavadora 0,20 · inodoro con cisterna 0,10 (L/s, agua fría).\n"
                    "- **Presión:** mínimo 100 kPa en grifos y 150 kPa en fluxores y calentadores; "
                    "máximo 500 kPa.\n- **DB-HS5:** evacuación de aguas.")
    with n3:
        st.markdown("- **RITE (RD 1027/2007):** regulación por termostato, eficiencia energética y "
                    "mantenimiento de instalaciones térmicas.\n"
                    f"- **Estado a {NORMATIVA_FECHA}:** modificaciones del RITE y del CTE anunciadas; "
                    "comprueba el BOE antes de cada proyecto.")
    with n4:
        st.markdown("- **Reglamento ICT (RD 346/2011):** infraestructuras comunes de telecomunicaciones. "
                    "La red interior de datos (Cat6) se presupuesta por tomas.")
    with n5:
        st.markdown("- **Protección de datos:** la videovigilancia con fines de seguridad se rige por el "
                    "RGPD y la LOPDGDD; la AEPD publica una guía, fichas y un modelo de cartel.\n"
                    "- **Cartel informativo** de zona videovigilada visible antes de entrar en el "
                    "campo de visión de las cámaras.\n"
                    "- **Ley 5/2014 de Seguridad Privada:** si el sistema se conecta a una central "
                    "receptora de alarmas, la instalación y el mantenimiento los hace una empresa de "
                    "seguridad habilitada.\n"
                    "- **Almacenamiento:** el cálculo de la app es orientativo (H.265, grabación "
                    "continua); el consumo real depende del códec, la escena y la compresión.")
